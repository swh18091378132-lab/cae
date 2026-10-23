# -*- coding: utf-8 -*-
"""
test_m7.py -- 票 07 验收：强度校核报告（.docx 模板灌数）

运行 (在 apps/workbench 下): python tests/test_m7.py
"""

import os
import sys
import shutil
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
WORKBENCH = os.path.dirname(HERE)
CORE = os.path.join(WORKBENCH, "core")
sys.path.insert(0, CORE)

import db
import study
import run_study
import report_template
import report as report_mod
import hotspot_figure
from catia_driver import FakeCatiaDriver
from abaqus_runner import FakeAbaqusRunner


PARAMS = {"span": 5.0, "sweep_deg": 10.0, "c_root": 1.0, "c_tip": 0.6,
          "n_ribs": 3, "n_stringers": 4}


def _tmpdir(name):
    d = os.path.join(WORKBENCH, ".tmp_test", name)
    if os.path.isdir(d):
        shutil.rmtree(d)
    os.makedirs(d)
    return d


def _doc_text(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def test_template(fails):
    tmp = _tmpdir("m7tpl")
    path = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(path)
    if not os.path.isfile(path):
        fails.append("template not generated")
        return

    from docx import Document
    try:
        doc = Document(path)
    except Exception as exc:
        fails.append("template not openable: %s" % exc)
        return

    text = _doc_text(path)
    missing = [p for p in report_template.PLACEHOLDERS if "{{%s}}" % p not in text
               and "{{ %s }}" % p not in text]
    if missing:
        fails.append("template missing placeholders: %s" % missing)


def test_render(fails):
    tmp = _tmpdir("m7render")
    tpl = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(tpl)
    out = os.path.join(tmp, "r.docx")

    ctx = {
        "study_id": "STUDY-X",
        "study_name": "STUDY-X",
        "template_id": "wingbox_v1",
        "generated_at": "2026-01-01 00:00:00",
        "param": {"span": 5.0, "sweep_deg": 10.0, "c_root": 1.0, "c_tip": 0.6,
                  "n_ribs": 3, "n_stringers": 4},
        "model": {"n_bays_total": 6, "n_bays_valid": 6, "n_bays_invalid": 0,
                  "n_elems": 6, "n_nodes": 20, "n_plies": 2,
                  "total_thickness": 0.2},
        "summary": {"max_fi": "0.5", "min_ms": "1.0", "all_pass": "合格",
                    "n_elems": 6, "n_detail_rows": 36,
                    "n_bays_checked": 6, "n_bays_excluded": 0},
        "check": {"criteria": "tsai_wu", "safety_factor": "1.5",
                  "ms_threshold": "0", "allowables": "Xt=1500 MPa"},
        # P2 新增命名空间（票14 分析类型 / 票15 网格质量）：
        # 模板已声明这些占位符，手工构造的 ctx 也得给值，否则会被判"未解析"。
        "analysis": {"kind": "static", "label": "静力分析",
                     "detail": "静力分析（P1 默认）", "is_linear": "是"},
        "mesh_quality": {"enabled": False,
                         "summary": "未启用网格质量门（按 P1 默认网格策略）",
                         "warnings": "无", "n_warn_bays": 0,
                         "n_refined_bays": 0, "worst_jacobian": "—",
                         "max_aspect": "—", "max_warp": "—", "criteria": ""},
        "allowables": "Xt=1500 MPa",
        "param_table": {"span": 5.0, "n_ribs": 3},
        "detail_table": [
            {"elem": 5, "sp": 2, "criterion": "tsai_wu", "FI": 0.9,
             "MS": 0.11, "pass": True},
        ],
        "hotspot_table": [
            {"elem": 5, "sp": 2, "criterion": "tsai_wu", "FI": 0.9, "MS": 0.11},
            {"elem": 3, "sp": 1, "criterion": "hashin", "FI": 0.7, "MS": 0.43},
        ],
        "invalid_bay_table": ["NB-r1-s0", "NB-r1-s1"],
        "artifact_table": [
            {"kind": "check_result", "version": 1,
             "checksum": "abcdef123456", "uri": "check/check_summary.json"},
        ],
        "hotspot_figure": None,
        "conclusion": "TEST-CONCLUSION",
    }

    rr = report_mod.render_report(ctx, tpl, out)
    if not os.path.isfile(out):
        fails.append("report not written")
        return

    text = _doc_text(out)
    # 文本占位符替换
    if "STUDY-X" not in text:
        fails.append("study_id not replaced")
    if "TEST-CONCLUSION" not in text:
        fails.append("conclusion not replaced")
    if "0.5" not in text:
        fails.append("summary.max_fi not replaced")
    if "合格" not in text:
        fails.append("all_pass not replaced")

    # 表格渲染
    from docx import Document
    doc = Document(out)
    table_sizes = [(len(t.rows), len(t.columns)) for t in doc.tables]
    # 期望：param_table(1+2 行) / hotspot_table(1+2) / invalid(1+2)
    if not any(r == 3 for (r, c) in table_sizes):
        fails.append("expected a 3-row table, got %s" % table_sizes)
    if len(doc.tables) < 3:
        fails.append("expected >=3 tables, got %d" % len(doc.tables))

    # 无未解析占位符
    if rr.get("unresolved"):
        fails.append("unresolved placeholders remain: %s" % rr["unresolved"])


def test_unknown_placeholder(fails):
    tmp = _tmpdir("m7unknown")
    tpl = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(tpl)

    # 往模板里塞一个未知占位符
    from docx import Document
    doc = Document(tpl)
    doc.add_paragraph("未知字段：{{not.a.real.key}}")
    doc.save(tpl)

    out = os.path.join(tmp, "r.docx")
    ctx = report_mod.build_report_context(
        {"study_id": "S", "parameters": {}, "check": {"summary": {}}})
    rr = report_mod.render_report(ctx, tpl, out)

    if not any("not.a.real.key" in w for w in rr.get("warnings") or []):
        fails.append("unknown placeholder not warned: %s" % rr.get("warnings"))
    # 保留原样
    text = _doc_text(out)
    if "{{not.a.real.key}}" not in text:
        fails.append("unknown placeholder should be preserved verbatim")


def test_figure(fails):
    tmp = _tmpdir("m7fig")
    # 用 fake 全链拿到 result
    res, conn = _full_chain(tmp, "fig")
    conn.close()
    png = os.path.join(tmp, "hot.png")
    hotspot_figure.draw_hotspot_figure(res, png)
    if not os.path.isfile(png):
        fails.append("figure not written")
    else:
        if os.path.getsize(png) < 1000:
            fails.append("figure suspiciously small: %d" % os.path.getsize(png))


def _full_chain(tmp, tag, with_report=True):
    conn = db.connect(os.path.join(tmp, "db.sqlite"))
    study.create_study("S-" + tag, "M7 " + tag, conn=conn)
    res = run_study.run_study(
        "S-" + tag, PARAMS, "wingbox_v1",
        FakeCatiaDriver(os.path.join(tmp, "catparts")), conn,
        params_dir=os.path.join(tmp, "params"),
        topo_dir=os.path.join(tmp, "topo"),
        with_fe_model=True, fe_dir=os.path.join(tmp, "fe"),
        subdiv_per_bay=1,
        plies=[{"angle": 0.0, "thickness": 0.1},
               {"angle": 90.0, "thickness": 0.1}],
        abaqus=FakeAbaqusRunner(n_elems=6, n_plies=2),
        with_solve=True, job_name="job1",
        solve_dir=os.path.join(tmp, "solve"),
        with_check=True, check_dir=os.path.join(tmp, "check"),
        rules={"criteria": ["tsai_wu"]},
        with_report=with_report,
        report_dir=os.path.join(tmp, "report"))
    return res, conn


def test_full_chain_report(fails):
    tmp = _tmpdir("m7chain")
    res, conn = _full_chain(tmp, "chain")

    arts = res["artifacts"]
    if "report" not in arts:
        fails.append("report artifact missing")
    else:
        # 血缘 check_result -> report
        n = conn.execute(
            "SELECT COUNT(*) c FROM edges WHERE parent_id=? AND child_id=?",
            (arts["check_result"], arts["report"])).fetchone()["c"]
        if n != 1:
            fails.append("lineage check_result->report missing")

    rep = res.get("report") or {}
    path = rep.get("path")
    if not path or not os.path.isfile(path):
        fails.append("report file missing: %r" % path)
    else:
        if os.path.getsize(path) < 5000:
            fails.append("report too small: %d" % os.path.getsize(path))
        # 能重新打开
        from docx import Document
        try:
            Document(path)
        except Exception as exc:
            fails.append("report not openable: %s" % exc)
        # 含图片
        with zipfile.ZipFile(path) as z:
            imgs = [n for n in z.namelist() if n.startswith("word/media/")]
        if not imgs:
            fails.append("report contains no images")
        # 关键内容
        text = _doc_text(path)
        if "强度校核报告" not in text:
            fails.append("report missing title")
        if "S-chain" not in text:
            fails.append("report missing study id")
        if rep.get("unresolved"):
            fails.append("report has unresolved placeholders: %s"
                         % rep["unresolved"])

    conn.close()


def _chain_ctx():
    """造一个贴近 run_study 真实返回的 context（不跑全链，快）。"""
    return report_mod.build_report_context({
        "study_id": "S-CTX",
        "template": "wingbox_v1",
        "parameters": {"span": 12.34, "sweep_deg": 25.0, "c_root": 3.0,
                       "c_tip": 1.2, "n_ribs": 6, "n_stringers": 7},
        "topology": {"n_total": 30, "n_valid": 28, "n_invalid": 2,
                     "invalid_ids": ["NB-r0-s1"]},
        "fe": {"n_nodes": 517, "n_elems": 480, "n_plies": 8,
               "total_thickness": 0.8},
        "solve": {"job": "job1", "status": "COMPLETED"},
        "check": {
            "summary": {"max_fi": 0.8123456, "min_ms": 0.231, "all_pass": True,
                        "n_elems": 480, "n_detail_rows": 3,
                        "n_bays_checked": 28, "n_bays_excluded": 2,
                        "criteria": ["tsai_wu"], "safety_factor": 1.5,
                        "allowables": {"Xt": 1500, "Xc": 1200, "Yt": 50,
                                       "Yc": 200, "S": 70}},
            "hotspots": [{"elem": 101, "sp": 3, "criterion": "tsai_wu",
                          "FI": 0.8123456, "MS": 0.231}],
            "invalid_bays": ["NB-r0-s1"],
            "detail_rows": [
                {"elem": 101, "sp": 3, "criterion": "tsai_wu",
                 "FI": 0.8123456, "MS": 0.231, "pass": True},
                {"elem": 77, "sp": 1, "criterion": "hashin",
                 "FI": 0.51, "MS": 0.96, "pass": True},
                {"elem": 9, "sp": 2, "criterion": "max_stress",
                 "FI": 0.4, "MS": 1.5, "pass": True},
            ]},
    })


def test_placeholder_namespaces(fails):
    """
    票 07 指定的占位符语法：{{param.span}} / {{summary.max_fi}} /
    {{check.criteria}} / {{model.*}}。逐参数与模型信息必须能灌进去。
    """
    tmp = _tmpdir("m7ns")
    tpl = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(tpl)

    # 标准模板自身就要带上这些占位符（用户替换时按此对照）
    tpl_text = _doc_text(tpl)
    for key in ("{{param.span}}", "{{param.n_ribs}}", "{{model.n_elems}}",
                "{{model.n_bays_valid}}", "{{detail_table}}"):
        if key not in tpl_text:
            fails.append("标准模板缺少占位符 %s" % key)

    ctx = _chain_ctx()
    out = os.path.join(tmp, "r.docx")
    rr = report_mod.render_report(ctx, tpl, out)
    text = _doc_text(out)

    if "12.34" not in text:
        fails.append("{{param.span}} 未灌入（找不到 12.34）")
    if "480" not in text:
        fails.append("{{model.n_elems}} 未灌入（找不到 480）")
    if rr.get("unresolved"):
        fails.append("标准模板渲染后仍有未解析占位符: %s" % rr["unresolved"])


def test_render_signature_with_study_id(fails):
    """票 07 契约：render_report(study_id, context, template_path, out_path)。"""
    tmp = _tmpdir("m7sig")
    tpl = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(tpl)
    out = os.path.join(tmp, "r.docx")

    ctx = _chain_ctx()
    rr = report_mod.render_report("S-SIG", ctx, tpl, out)
    if not os.path.isfile(out):
        fails.append("study_id 签名下报告未生成")
        return
    text = _doc_text(out)
    if "S-SIG" not in text:
        fails.append("study_id 实参未写入报告")
    if rr.get("path") != out:
        fails.append("study_id 签名下返回路径不符")


def test_detail_table(fails):
    """校核明细表：行数 = 明细行数（含表头 +1）。"""
    tmp = _tmpdir("m7detail")
    tpl = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(tpl)
    out = os.path.join(tmp, "r.docx")

    ctx = _chain_ctx()
    n = len(ctx["detail_table"])
    if n != 3:
        fails.append("detail_table 应有 3 行，实际 %d" % n)
    report_mod.render_report(ctx, tpl, out)

    from docx import Document
    doc = Document(out)
    found = None
    for t in doc.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if "单元" in hdr and "准则" in hdr and "铺层" in hdr and "FI" in hdr:
            if "序号" not in hdr:      # 区分危险部位表（带序号）
                found = t
                break
    if found is None:
        fails.append("未找到校核明细表（表头应含 单元/铺层/准则/FI）")
    elif len(found.rows) - 1 != n:
        fails.append("校核明细表行数 %d != 明细行数 %d"
                     % (len(found.rows) - 1, n))


def main():
    fails = []
    print("1. template ...")
    test_template(fails)
    print("2. render ...")
    test_render(fails)
    print("3. unknown placeholder ...")
    test_unknown_placeholder(fails)
    print("4. chain + report ...")
    test_full_chain_report(fails)
    print("5. figure ...")
    test_figure(fails)
    print("6. placeholder namespaces (param./model./detail) ...")
    test_placeholder_namespaces(fails)
    print("7. render_report(study_id, ...) signature ...")
    test_render_signature_with_study_id(fails)
    print("8. detail table ...")
    test_detail_table(fails)

    if fails:
        print("FAILED:")
        for f in fails:
            print("  - " + f)
        return 1
    print("M7 (票房07) SELF-TEST: ALL PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())

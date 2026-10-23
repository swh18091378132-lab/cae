# -*- coding: utf-8 -*-
"""
test_m13.py -- 票13 验收：机身筒段部件族（筒段拓扑 + 端到端闭环）

验收点（对应 issue #13）：
  1. 模板列表可选到机身筒段，参数契约含框位、桁条数、轴向长度、半径、蒙皮开口位置
  2. 自然格由「框 × 桁条」围合生成，数量与布置规则一致（闭口含接缝）
  3. 蒙皮开口造成的不连续格被正确标记并说明原因（cutout）
  4. 给定合法参数集，跑通建模→拓扑→FE→求解→抽取→校核→报告全链（fake 适配器）
  5. 血缘链完整，报告可由该 Study 生成
  6. 参数越界/未知参数按契约被拒绝
  7. Python 3.8 兼容

设计要点：
  * 筒段类是票11 拓扑策略化的**第二个**模式（cylinder_section）：框沿轴向等分（对应翼面
    的"肋"），桁条沿周向等分（对应翼面的"长桁"），闭口故在周向接缝 xi=1 补一根桁条。
    自然格 = (n_frames-1) × n_stringers。
  * 蒙皮开口位置在模板里用**参数**表达（hole_axial_frac / hole_circ_frac / hole_enabled），
    映射成拓扑层的 holes = [{'cy': 沿轴向分数, 'cxi': 沿周向分数}]。两类部件族因此共用
    同一套不连续格判据（ADR-0008），本票只验证映射与结果。
  * 报告：既有标准模板是翼面族口径（含 param.span 等），筒段 Study 渲染时这些键不可解析，
    渲染器按既有约定**保留原样并告警**（不静默丢弃）——本票断言"未解析的只剩翼面族专属键"，
    并用一份筒段口径模板断言可产出零残留占位符的报告。

运行 (在 apps/workbench 下): python tests/test_m13.py
"""

import os
import sys
import json
import shutil

HERE = os.path.dirname(os.path.abspath(__file__))
WORKBENCH = os.path.dirname(HERE)
CORE = os.path.join(WORKBENCH, "core")
sys.path.insert(0, CORE)
sys.path.insert(0, WORKBENCH)

import db
import study
import topology
import run_study
import template_registry
import report as report_mod
from catia_driver import FakeCatiaDriver
from abaqus_runner import FakeAbaqusRunner


TEMPLATE_ID = "fuselage_section_v1"

# 合法参数集：4 框（含两端）× 8 桁条（+接缝）→ 3 × 8 = 24 个自然格；
# 开口落 y=3.0m（轴长 6m 的一半）、xi=0.5625（第 5 段周向格内）→ 恰好 1 个不连续格。
PARAMS = {
    "length": 6.0,
    "radius": 1.5,
    "n_frames": 4,
    "n_stringers": 8,
    "hole_axial_frac": 0.5,
    "hole_circ_frac": 0.5625,
    "hole_enabled": 1,
}

# 与翼面族一一对应的专属占位符（标准报告模板里含这些键；筒段 Study 无对应参数）
WING_ONLY_PLACEHOLDERS = ("param.span", "param.sweep_deg", "param.c_root",
                          "param.c_tip", "param.n_ribs")

PLIES = [{"angle": 0.0, "thickness": 0.1},
         {"angle": 90.0, "thickness": 0.1}]

RULES = {"criteria": ["tsai_wu", "hashin"], "safety_factor": 1.0}


# ---------------------------------------------------------------------------
# 工具
# ---------------------------------------------------------------------------

def _fresh(name):
    d = os.path.join(WORKBENCH, ".tmp_test", "m13", name)
    if os.path.isdir(d):
        shutil.rmtree(d)
    os.makedirs(d)
    return d


def _near(a, b, tol=1e-9):
    return abs(float(a) - float(b)) <= tol


def holes_from_params(params):
    """
    模板参数 → 拓扑层 holes（票13 的参数契约到通用口径的映射）。

    hole_axial_frac 是**沿轴向的分数**（0=前端框，1=后端框），hole_circ_frac 是
    **沿周向的分数**（0/1=接缝）；两者正好等于拓扑层的 cy / cxi 口径，故是恒等映射。
    这两者与部件族无关——翼面类的 cy/cxi 是同一定义（沿 y / 沿 xi 的分数）。
    """
    if not int(params.get("hole_enabled", 1)):
        return []
    return [{"cy": float(params["hole_axial_frac"]),
             "cxi": float(params["hole_circ_frac"])}]


def _ids(model):
    return [b.id for b in model.invalid_bays]


def _doc_text(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _expected_n_bays(params):
    """布置规则：自然格数 = (框数-1) × 桁条数（闭口接缝那一段也是独立的一格）。"""
    return (int(params["n_frames"]) - 1) * int(params["n_stringers"])


# ---------------------------------------------------------------------------
# 1) 模板可选到 + 参数契约
# ---------------------------------------------------------------------------

def test_template_contract(fails):
    """模板列表能选到机身筒段；契约含框位/桁条数/轴向长度/半径/开口位置。"""
    listed = template_registry.list_templates()
    if TEMPLATE_ID not in listed:
        fails.append("模板列表里没有 %s（现有 %r）" % (TEMPLATE_ID, listed))

    # 工作台服务层（UI 取模板列表的唯一入口）同样要能看到
    import workbench_service as ws
    root = _fresh("svc_root")
    svc = ws.WorkbenchService(root, mode="fake")
    try:
        if TEMPLATE_ID not in svc.list_templates():
            fails.append("服务层 list_templates() 选不到机身筒段")
        if svc.get_template(TEMPLATE_ID).id != TEMPLATE_ID:
            fails.append("服务层 get_template 返回的模板 id 不符")
    finally:
        svc.close()

    tpl = template_registry.load(TEMPLATE_ID)
    c = tpl.contract()
    for name, label in (("n_frames", "框位"),
                        ("n_stringers", "桁条数"),
                        ("length", "轴向长度"),
                        ("radius", "半径"),
                        ("hole_axial_frac", "开口轴向位置"),
                        ("hole_circ_frac", "开口周向位置")):
        if name not in c:
            fails.append("参数契约缺少 %s（%s）" % (name, label))
            continue
        spec = c[name]
        for k in ("type", "min", "max", "default"):
            if k not in spec:
                fails.append("参数 %s 的契约缺少 %s" % (name, k))

    # 契约里的默认值本身必须能通过校验（模板自洽）
    try:
        ok = tpl.validate({})
    except template_registry.ParamError as exc:
        fails.append("模板默认值不满足自身契约: %s" % exc)
    else:
        if not (0.0 <= ok["hole_axial_frac"] <= 1.0):
            fails.append("开口轴向位置默认值不在 0..1: %r" % ok["hole_axial_frac"])
        if not (0.0 <= ok["hole_circ_frac"] <= 1.0):
            fails.append("开口周向位置默认值不在 0..1: %r" % ok["hole_circ_frac"])

    if tpl.envelope.get("constraints") and not tpl.envelope.get("note"):
        fails.append("参数化包络缺少 note")

    # 拓扑声明必须分派到筒段模式，且取参与契约一致
    rule = topology.topology_rule(template=tpl)
    if rule.get("mode") != topology.MODE_CYLINDER_SECTION:
        fails.append("模板拓扑模式应为 cylinder_section，实际 %r" % rule.get("mode"))
    pp = rule.get("params") or {}
    for k, expect in (("axial", "length"), ("n_frames", "n_frames"),
                      ("n_stringers", "n_stringers"), ("radius", "radius")):
        if pp.get(k) != expect:
            fails.append("topology.params[%s] 应映射到 %r，实际 %r"
                         % (k, expect, pp.get(k)))


# ---------------------------------------------------------------------------
# 2) 自然格 = 框 × 桁条围合，数量与布置规则一致
# ---------------------------------------------------------------------------

def test_bays_from_frames_and_stringers(fails):
    """框沿轴向等分、桁条沿周向等分（闭口含接缝）；格数 = (框-1) × 桁条。"""
    params = dict(PARAMS)
    params["hole_enabled"] = 0
    model = topology.build_topology(TEMPLATE_ID, params)

    if model.mode != topology.MODE_CYLINDER_SECTION:
        fails.append("模型未记录 cylinder_section 模式: %r" % model.mode)

    n_frames = int(params["n_frames"])
    n_str = int(params["n_stringers"])
    length = float(params["length"])
    radius = float(params["radius"])

    # 框（= 翼面族的"肋"）：n_frames 根含两端，沿轴向等分
    if len(model.ribs) != n_frames:
        fails.append("框数应为 %d，实际 %d" % (n_frames, len(model.ribs)))
    else:
        for k, rib in enumerate(model.ribs):
            want_y = length * k / float(n_frames - 1)
            if not _near(rib["y"], want_y):
                fails.append("第 %d 根框未等分轴向：y=%.6f，期望 %.6f"
                             % (k, rib["y"], want_y))
            if not (_near(rib["xi0"], 0.0) and _near(rib["xi1"], 1.0)):
                fails.append("框 %d 未覆盖整圈周向: [%r, %r]"
                             % (k, rib["xi0"], rib["xi1"]))

    # 桁条：n_stringers 根等分周向 + 闭口接缝 1 根
    if len(model.stringers) != n_str + 1:
        fails.append("桁条数应为 %d（含接缝），实际 %d"
                     % (n_str + 1, len(model.stringers)))
    else:
        for k in range(n_str):
            want_xi = k / float(n_str)
            if not _near(model.stringers[k]["xi"], want_xi):
                fails.append("第 %d 根桁条未等分周向：xi=%.6f，期望 %.6f"
                             % (k, model.stringers[k]["xi"], want_xi))
        if not _near(model.stringers[-1]["xi"], 1.0):
            fails.append("闭口接缝桁条应落在 xi=1.0，实际 %r"
                         % model.stringers[-1]["xi"])
        for k, s in enumerate(model.stringers):
            if not (_near(s["y0"], 0.0) and _near(s["y1"], length)):
                fails.append("桁条 %d 未覆盖全轴向: y[%r, %r]"
                             % (k, s["y0"], s["y1"]))

    # 自然格数量
    expect = _expected_n_bays(params)
    if len(model.bays) != expect:
        fails.append("自然格数应为 %d，实际 %d" % (expect, len(model.bays)))
    if len(model.invalid_bays) != 0:
        fails.append("无开口时不应有不连续格: %r" % _ids(model))

    # 每个格的四边确实由「相邻两框 × 相邻两桁条」围合
    for bay in model.bays:
        r0, r1 = model.ribs[bay.rib_j], model.ribs[bay.rib_j + 1]
        s0, s1 = model.stringers[bay.str_i], model.stringers[bay.str_i + 1]
        if not (_near(bay.y0, r0["y"]) and _near(bay.y1, r1["y"])):
            fails.append("格 %s 的轴向边界不是相邻框" % bay.id)
        if not (_near(bay.xi0, s0["xi"]) and _near(bay.xi1, s1["xi"])):
            fails.append("格 %s 的周向边界不是相邻桁条" % bay.id)

    # 所有格拼起来 = 整个筒段蒙皮（无空缺、无重叠）
    area = sum((b.y1 - b.y0) * (b.xi1 - b.xi0) for b in model.bays)
    if not _near(area, length * 1.0, 1e-9):
        fails.append("自然格总面积 %.9f != 筒段蒙皮面积 %.9f" % (area, length))

    # 物理尺度：xi 是周向 0..1，换算成展开周长 2πr（供 FE / 网格质量门用）
    ext = model.extents()
    if not _near(ext["xi_scale"], 2.0 * 3.141592653589793 * radius, 1e-9):
        fails.append("xi_scale 应为展开周长 2πr=%.6f，实际 %.6f"
                     % (2.0 * 3.141592653589793 * radius, ext["xi_scale"]))
    if not _near(ext["y_span"], length):
        fails.append("y_span 应为轴向长度 %r，实际 %r" % (length, ext["y_span"]))
    if not _near(ext["xi_span"], 2.0 * 3.141592653589793 * radius, 1e-9):
        fails.append("xi_span 应为展开周长，实际 %r" % ext["xi_span"])

    # 布置规则随参数缩放：5 框 × 12 桁条 → 4 × 12 = 48 格
    p2 = dict(params)
    p2["n_frames"] = 5
    p2["n_stringers"] = 12
    m2 = topology.build_topology(TEMPLATE_ID, p2)
    if len(m2.bays) != _expected_n_bays(p2):
        fails.append("5 框 × 12 桁条应有 %d 格，实际 %d"
                     % (_expected_n_bays(p2), len(m2.bays)))
    if len(m2.invalid_bays) != 0:
        fails.append("5 框 × 12 桁条（无开口）不应有不连续格: %r" % _ids(m2))

    # 稳定 ID：同参数集两次生成一致且唯一
    again = topology.build_topology(TEMPLATE_ID, params)
    ids_a = [b.id for b in model.bays]
    ids_b = [b.id for b in again.bays]
    if ids_a != ids_b:
        fails.append("同参数集两次生成的格 ID 不一致")
    if len(set(ids_b)) != len(ids_b):
        fails.append("自然格 ID 不唯一")

    # 序列化（拓扑产物落盘用）
    data = json.loads(model.to_json())
    if data.get("topology_mode") != topology.MODE_CYLINDER_SECTION:
        fails.append("拓扑 JSON 未记录 cylinder_section")
    if data["n_bays_total"] != expect or data["n_bays_valid"] != expect:
        fails.append("拓扑 JSON 计数不符: %r"
                     % {k: data[k] for k in ("n_bays_total", "n_bays_valid")})


# ---------------------------------------------------------------------------
# 3) 蒙皮开口造成的不连续格被正确标记，原因正确
# ---------------------------------------------------------------------------

def test_cutout_invalid_bays(fails):
    """开口 → 包含它的那个格被标为不连续，原因是 cutout 且说明可读。"""
    model = topology.build_topology(TEMPLATE_ID, PARAMS,
                                   holes=holes_from_params(PARAMS))

    # 开口：y = 0.5 × 6m = 3.0m（落在框 1 与框 2 之间，j=1）
    #       xi = 0.5625（落在桁条 4 与桁条 5 之间，i=4）
    if len(model.invalid_bays) != 1:
        fails.append("开口应造成 1 个不连续格，实际 %d（%r）"
                     % (len(model.invalid_bays), _ids(model)))
        return
    bay = model.invalid_bays[0]
    if bay.id != "NB-r1-s4":
        fails.append("被标不连续的格应是「框1-2 × 桁条4-5」= NB-r1-s4，实际 %s" % bay.id)
    if bay.valid:
        fails.append("不连续格的 valid 标志未置 False")

    types = [r["type"] for r in bay.reasons]
    if types != [topology.REASON_CUTOUT]:
        fails.append("不连续原因应恰好是 cutout，实际 %r" % types)
    detail = " ".join(r.get("detail", "") for r in bay.reasons)
    if "cutout" not in detail:
        fails.append("不连续原因缺少可读说明: %r" % detail)
    # 说明里要带上开口的实际坐标，便于工程师定位
    if "3.000" not in detail or "0.5625" not in detail:
        fails.append("原因说明未给出开口坐标: %r" % detail)

    # 其余格照常有效
    if len(model.valid_bays) != _expected_n_bays(PARAMS) - 1:
        fails.append("有效格数应为 %d，实际 %d"
                     % (_expected_n_bays(PARAMS) - 1, len(model.valid_bays)))

    # 开口落在框线上（不是格内部）→ 不造成格不连续
    p = dict(PARAMS)
    p["hole_axial_frac"] = 1.0 / 3.0     # y = 2.0m = 框 1 的位置
    m2 = topology.build_topology(TEMPLATE_ID, p, holes=holes_from_params(p))
    if len(m2.invalid_bays) != 0:
        fails.append("开口落在框线上不应判不连续格，实际 %r" % _ids(m2))

    # 开口落在接缝那一段周向格里（最后一根桁条与接缝之间）→ 照常标记
    p3 = dict(PARAMS)
    p3["hole_circ_frac"] = 0.99
    m3 = topology.build_topology(TEMPLATE_ID, p3, holes=holes_from_params(p3))
    if _ids(m3) != ["NB-r1-s7"]:
        fails.append("接缝段的开口应标记 NB-r1-s7，实际 %r" % _ids(m3))

    # 关闭开口 → 无不连续格
    p4 = dict(PARAMS)
    p4["hole_enabled"] = 0
    m4 = topology.build_topology(TEMPLATE_ID, p4, holes=holes_from_params(p4))
    if _ids(m4) != []:
        fails.append("hole_enabled=0 时不应有不连续格: %r" % _ids(m4))

    # 多开口：每个开口各自标记自己所在的格
    m5 = topology.build_topology(
        TEMPLATE_ID, PARAMS,
        holes=[{"cy": 0.5, "cxi": 0.5625},      # j=1, i=4
               {"cy": 0.9, "cxi": 0.0625}])     # j=2, i=0
    if sorted(_ids(m5)) != ["NB-r1-s4", "NB-r2-s0"]:
        fails.append("两个开口应各标记一格，实际 %r" % sorted(_ids(m5)))


# ---------------------------------------------------------------------------
# 4) 端到端全链（fake 适配器）
# ---------------------------------------------------------------------------

def _full_chain(tmp, tag, report_template_path=None, src_params=None):
    """跑一次 fake 全链，返回 (result, conn, study_id)。"""
    params = dict(src_params or PARAMS)
    conn = db.connect(os.path.join(tmp, "m13_%s.db" % tag))
    sid = "S-M13-" + tag
    study.create_study(sid, "M13 机身筒段 " + tag, conn=conn)

    opts = run_study.RunOptions(
        params_dir=os.path.join(tmp, tag, "params"),
        topo_dir=os.path.join(tmp, tag, "topo"),
        with_topology=True,
        with_fe_model=True,
        fe_dir=os.path.join(tmp, tag, "fe"),
        subdiv_per_bay=1,
        plies=PLIES,
        holes=holes_from_params(params),
        abaqus=FakeAbaqusRunner(n_elems=_expected_n_bays(params), n_plies=2),
        with_solve=True, job_name="job1",
        solve_dir=os.path.join(tmp, tag, "solve"),
        with_check=True, check_dir=os.path.join(tmp, tag, "check"),
        rules=RULES,
        with_report=True, report_dir=os.path.join(tmp, tag, "report"))
    if report_template_path:
        opts.set("report_template_path", report_template_path)

    res = run_study.run_study(
        sid, params, TEMPLATE_ID,
        FakeCatiaDriver(os.path.join(tmp, tag, "catparts")), conn, opts)
    return res, conn, sid


def test_full_chain(fails):
    """合法参数集跑通 建模→拓扑→FE→求解→抽取→校核→报告 全链。"""
    tmp = _fresh("chain")
    res, conn, sid = _full_chain(tmp, "chain")
    try:
        arts = res.get("artifacts") or {}
        for kind in ("parameter_set", "catpart", "topology", "mesh", "inp",
                     "odb", "extraction", "check_rule", "check_result", "report"):
            if kind not in arts:
                fails.append("全链缺少产物 %s" % kind)

        # 数模确实是按机身筒段模板再生的
        if res.get("template") != TEMPLATE_ID:
            fails.append("result.template 应为 %s，实际 %r"
                         % (TEMPLATE_ID, res.get("template")))
        if not os.path.isfile(res.get("catpart_path") or ""):
            fails.append("参数化数模文件未产出: %r" % res.get("catpart_path"))
        if not os.path.isfile(res["fe"].get("inp_path") or ""):
            fails.append("FE 输入文件未产出")

        # 拓扑：24 个候选格，开口造成 1 个不连续格
        topo = res.get("topology") or {}
        expect = _expected_n_bays(PARAMS)
        if topo.get("n_total") != expect:
            fails.append("全链拓扑候选格数应为 %d，实际 %r" % (expect, topo.get("n_total")))
        if topo.get("n_invalid") != 1:
            fails.append("全链拓扑不连续格数应为 1，实际 %r" % topo.get("n_invalid"))
        if topo.get("invalid_ids") != ["NB-r1-s4"]:
            fails.append("全链拓扑不连续格 ID 不符: %r" % topo.get("invalid_ids"))

        # FE：网格按自然格划分，单元数 = 格数 × subdiv²（subdiv_per_bay=1）
        fe = res.get("fe") or {}
        if fe.get("n_elems") != expect:
            fails.append("单元数应为 %d，实际 %r" % (expect, fe.get("n_elems")))
        if fe.get("n_plies") != len(PLIES):
            fails.append("铺层数应为 %d，实际 %r" % (len(PLIES), fe.get("n_plies")))
        if not os.path.isfile(fe.get("inp_path") or ""):
            fails.append("inp 路径无效: %r" % fe.get("inp_path"))
        else:
            with open(fe["inp_path"], "r", encoding="utf-8") as fh:
                inp = fh.read()
            for token in ("*HEADING", "*NODE", "*ELEMENT, TYPE=S4R",
                          "*MATERIAL", "*SHELL SECTION", "*STEP", "*END STEP"):
                if token not in inp:
                    fails.append("inp 缺少 %r" % token)
            if TEMPLATE_ID not in inp:
                fails.append("inp 未记录模板族 %s（可追溯）" % TEMPLATE_ID)

        # 求解 + 抽取
        solve = res.get("solve") or {}
        if solve.get("status") != "COMPLETED":
            fails.append("求解未完成: %r" % solve.get("status"))
        ext = res.get("extraction") or {}
        if not os.path.isfile(ext.get("csv_path") or ""):
            fails.append("抽取 CSV 未产出: %r" % ext.get("csv_path"))

        # 校核：不连续格被排除出自然格级，单元级照常
        chk = res.get("check") or {}
        summ = chk.get("summary") or {}
        if summ.get("n_detail_rows", 0) <= 0:
            fails.append("校核无明细行")
        if summ.get("n_elems", 0) != expect:
            fails.append("校核覆盖单元数应为 %d，实际 %r"
                         % (expect, summ.get("n_elems")))
        if summ.get("n_bays_excluded") != 1:
            fails.append("应排除 1 个不连续格，实际 %r" % summ.get("n_bays_excluded"))
        if summ.get("n_bays_checked") != expect - 1:
            fails.append("自然格级应校核 %d 格，实际 %r"
                         % (expect - 1, summ.get("n_bays_checked")))
        if "NB-r1-s4" in (chk.get("bay_best") or {}):
            fails.append("不连续格 NB-r1-s4 不应出现在自然格级结果里")
        if not chk.get("hotspots"):
            fails.append("未定位危险部位")
        if chk.get("invalid_bays") != ["NB-r1-s4"]:
            fails.append("校核结果未列出不连续格: %r" % chk.get("invalid_bays"))

        # 报告
        rep = res.get("report") or {}
        if not rep.get("path") or not os.path.isfile(rep["path"]):
            fails.append("报告未生成: %r" % rep.get("path"))
        elif os.path.getsize(rep["path"]) < 5000:
            fails.append("报告过小: %d 字节" % os.path.getsize(rep["path"]))
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# 5) 血缘链完整 + 报告可由该 Study 生成
# ---------------------------------------------------------------------------

def test_lineage_and_report(fails):
    """血缘链完整；报告可由该 Study 生成并携带产物指纹。"""
    tmp = _fresh("lineage")
    tpl_path = os.path.join(tmp, "fuselage_report.docx")
    try:
        _build_fuselage_report_template(tpl_path)
    except ImportError:
        print("  [SKIP] 无 python-docx，跳过报告内容断言")
        tpl_path = None

    res, conn, sid = _full_chain(tmp, "lineage", report_template_path=tpl_path)
    try:
        arts = res.get("artifacts") or {}
        chain = [("parameter_set", "catpart"), ("catpart", "topology"),
                 ("topology", "mesh"), ("mesh", "inp"), ("inp", "odb"),
                 ("odb", "extraction"), ("extraction", "check_result"),
                 ("check_rule", "check_result"), ("check_result", "report")]
        for parent, child in chain:
            if parent not in arts or child not in arts:
                fails.append("血缘链缺少产物 %s/%s" % (parent, child))
                continue
            n = conn.execute(
                "SELECT COUNT(*) c FROM edges WHERE parent_id=? AND child_id=?",
                (arts[parent], arts[child])).fetchone()["c"]
            if n != 1:
                fails.append("血缘边 %s->%s 缺失" % (parent, child))

        # 参数集产物里落的是校验后的机身筒段参数
        row = conn.execute("SELECT params_json FROM artifacts WHERE id=?",
                           (arts["parameter_set"],)).fetchone()
        saved = json.loads(row["params_json"])
        for k in ("n_frames", "n_stringers", "length", "radius",
                  "hole_axial_frac", "hole_circ_frac"):
            if k not in saved:
                fails.append("参数集产物缺少 %s" % k)

        # 报告产物：kind=report，挂在该 Study 下，可由 Study 取回
        rep = res.get("report") or {}
        path = rep.get("path")
        if not path or not os.path.isfile(path):
            fails.append("报告文件缺失: %r" % path)
        else:
            row = conn.execute(
                "SELECT study_id, kind, version, content_uri FROM artifacts WHERE id=?",
                (arts["report"],)).fetchone()
            if row["study_id"] != sid or row["kind"] != "report":
                fails.append("报告产物登记信息不符: %r" % dict(row))
            if os.path.abspath(row["content_uri"]) != os.path.abspath(path):
                fails.append("报告产物的 content_uri 与实际文件不一致")
            if not rep.get("figure") or not os.path.isfile(rep["figure"]):
                fails.append("危险部位图未生成: %r" % rep.get("figure"))

            # 标准模板里翼面族专属占位符对筒段 Study 不可解析：
            # 按既有约定保留原样并告警，绝不静默丢弃；其余字段必须全部解析成功。
            unresolved = list(rep.get("unresolved") or [])
            bad = [u for u in unresolved
                   if u.replace("{{", "").replace("}}", "").replace(" ", "")
                   not in WING_ONLY_PLACEHOLDERS]
            if bad:
                fails.append("报告里有非翼面族专属的未解析占位符: %r" % bad)
            text = _doc_text(path)
            if "强度校核报告" not in text:
                fails.append("报告缺少标题")
            if sid not in text:
                fails.append("报告未写入 Study id")
            if TEMPLATE_ID not in text:
                fails.append("报告未写入模板 id")

            # 自定义筒段口径模板：零残留占位符，参数与指纹全部灌入
            if tpl_path:
                out2 = os.path.join(tmp, "lineage", "fuselage_report_out.docx")
                ctx = report_mod.build_report_context(res)
                rr = report_mod.render_report(sid, ctx, tpl_path, out2)
                if rr.get("unresolved"):
                    fails.append("筒段口径模板仍有未解析占位符: %r" % rr["unresolved"])
                t2 = _doc_text(out2)
                if "机身筒段" not in t2:
                    fails.append("筒段口径报告缺少机身筒段标识")
                if "轴长=6.0 m" not in t2:
                    fails.append("筒段口径报告未灌入轴向长度（param.length）")
                if "框位=4" not in t2 or "桁条数=8" not in t2:
                    fails.append("筒段口径报告未灌入框位/桁条数")
                if "不连续 1" not in t2:
                    fails.append("筒段口径报告未灌入不连续格数")
                # 可追溯性：报告里带 check_result 产物指纹
                row = conn.execute("SELECT checksum FROM artifacts WHERE id=?",
                                   (arts["check_result"],)).fetchone()
                if row["checksum"][:12] not in t2:
                    fails.append("报告未携带 check_result 产物指纹（可追溯性）")

        # 报告可由该 Study 取回（服务层口径）
        import workbench_service as ws
        svc = ws.WorkbenchService(tmp, db_path=os.path.join(tmp, "m13_lineage.db"),
                                  mode="fake")
        try:
            got = svc.get_report_path(sid)
            if not got or not os.path.isfile(got):
                fails.append("服务层 get_report_path 取不回该 Study 的报告: %r" % got)
            lin = svc.get_lineage(sid)
            if len(lin.get("nodes") or []) < 10:
                fails.append("血缘节点数偏少: %d" % len(lin.get("nodes") or []))
        finally:
            svc.close()
    finally:
        conn.close()


def _build_fuselage_report_template(path):
    """造一份筒段口径的报告模板（模拟工程师按项目规范替换模板，票07 已允许）。"""
    from docx import Document
    doc = Document()
    doc.add_heading("机身筒段强度校核报告", level=0)
    doc.add_paragraph("分析任务：{{study_id}}    模板：{{template_id}}")
    doc.add_paragraph("轴长={{param.length}} m    半径={{param.radius}} m")
    doc.add_paragraph("框位={{param.n_frames}}    桁条数={{param.n_stringers}}")
    doc.add_paragraph("自然格：候选 {{model.n_bays_total}}，不连续 {{model.n_bays_invalid}}")
    doc.add_paragraph("单元 {{model.n_elems}}，网格节点 {{model.n_nodes}}")
    doc.add_paragraph("总体判定：{{summary.all_pass}}    最大 FI：{{summary.max_fi}}")
    doc.add_paragraph("{{param_table}}")
    doc.add_paragraph("{{detail_table}}")
    doc.add_paragraph("{{invalid_bay_table}}")
    doc.add_paragraph("{{artifact_table}}")
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 6) 参数越界 / 未知参数被拒绝
# ---------------------------------------------------------------------------

def test_param_rejection(fails):
    """参数越界、未知参数、包络越界都按契约被拒绝，且不落半成品产物。"""
    tpl = template_registry.load(TEMPLATE_ID)

    def _expect_reject(params, why):
        try:
            tpl.validate(params)
        except template_registry.ParamError:
            return
        fails.append("%s 未被拒绝: %r" % (why, params))

    # 未知参数
    bad = dict(PARAMS)
    bad["wingspan"] = 10.0
    _expect_reject(bad, "未知参数")
    # 越界：轴向长度
    bad = dict(PARAMS)
    bad["length"] = 999.0
    _expect_reject(bad, "轴向长度超上限")
    bad = dict(PARAMS)
    bad["length"] = 0.1
    _expect_reject(bad, "轴向长度低于下限")
    # 越界：半径 / 框位 / 桁条数
    bad = dict(PARAMS)
    bad["radius"] = 100.0
    _expect_reject(bad, "半径超上限")
    bad = dict(PARAMS)
    bad["n_frames"] = 1
    _expect_reject(bad, "框位低于下限（至少 2 根才能围出格）")
    bad = dict(PARAMS)
    bad["n_stringers"] = 2
    _expect_reject(bad, "桁条数低于下限")
    # 越界：开口位置（必须是 0..1 的分数）
    bad = dict(PARAMS)
    bad["hole_axial_frac"] = 1.5
    _expect_reject(bad, "开口轴向分数超 1")
    bad = dict(PARAMS)
    bad["hole_circ_frac"] = -0.1
    _expect_reject(bad, "开口周向分数小于 0")
    # 类型错误：框位必须是 int
    bad = dict(PARAMS)
    bad["n_frames"] = 4.5
    _expect_reject(bad, "框位非整数")
    # 包络（跨参数）：框密到超过工艺下限
    bad = dict(PARAMS)
    bad["length"] = 1.0
    bad["n_frames"] = 40
    _expect_reject(bad, "框间距低于包络下限（length/(n_frames-1) < 0.3）")
    # 包络（跨参数）：桁条周向间距过密
    bad = dict(PARAMS)
    bad["radius"] = 0.5
    bad["n_stringers"] = 60
    _expect_reject(bad, "桁条周向间距低于包络下限")

    # 合法参数仍应通过，且默认值补齐
    ok = tpl.validate(PARAMS)
    if ok["n_frames"] != PARAMS["n_frames"]:
        fails.append("合法参数校验后被篡改: %r" % ok)

    # 编排入口：越界参数直接抛，不产生任何产物
    tmp = _fresh("reject")
    conn = db.connect(os.path.join(tmp, "reject.db"))
    try:
        sid = "S-M13-reject"
        study.create_study(sid, "M13 拒绝测试", conn=conn)
        bad = dict(PARAMS)
        bad["length"] = 999.0
        raised = False
        try:
            run_study.run_study(
                sid, bad, TEMPLATE_ID,
                FakeCatiaDriver(os.path.join(tmp, "catparts")), conn,
                run_study.RunOptions(params_dir=os.path.join(tmp, "params")))
        except template_registry.ParamError:
            raised = True
        if not raised:
            fails.append("run_study 未拒绝越界参数")
        n = conn.execute("SELECT COUNT(*) c FROM artifacts WHERE study_id=?",
                         (sid,)).fetchone()["c"]
        if n != 0:
            fails.append("参数被拒绝时不应产生任何产物，实际 %d 个" % n)
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# 7) Python 3.8 兼容
# ---------------------------------------------------------------------------

def test_py38(fails):
    """本票涉及的模块不得使用 3.9+ 语法/标准库（目标机 Win7 + CPython 3.8）。"""
    import compat_py38
    targets = [
        os.path.join(CORE, "topology.py"),
        os.path.join(CORE, "template_registry.py"),
        os.path.join(CORE, "run_study.py"),
        os.path.join(HERE, "test_m13.py"),
    ]
    bad = compat_py38.check_paths(targets)
    if bad:
        fails.append("Python 3.8 兼容性检查未通过:\n%s"
                     % compat_py38.format_violations(bad))

    # 模板 JSON 必须能被 3.8 环境下的标准库读通（无注释/无尾逗号）
    path = os.path.join(os.path.dirname(CORE), "..", "..", "data", "templates",
                        TEMPLATE_ID + ".json")
    with open(os.path.abspath(path), "r", encoding="utf-8") as fh:
        data = json.load(fh)
    if data.get("id") != TEMPLATE_ID:
        fails.append("模板 JSON 的 id 字段不符: %r" % data.get("id"))


# ---------------------------------------------------------------------------

def main():
    fails = []
    print("1. 模板可选到 + 参数契约（框位/桁条数/开口位置） ...")
    test_template_contract(fails)
    print("2. 自然格由「框 × 桁条」围合，数量与布置规则一致 ...")
    test_bays_from_frames_and_stringers(fails)
    print("3. 蒙皮开口 → 不连续格被标记且原因正确 ...")
    test_cutout_invalid_bays(fails)
    print("4. 全链（建模→拓扑→FE→求解→抽取→校核→报告） ...")
    test_full_chain(fails)
    print("5. 血缘链 + 报告可由该 Study 生成 ...")
    test_lineage_and_report(fails)
    print("6. 参数越界/未知参数被拒绝 ...")
    test_param_rejection(fails)
    print("7. Python 3.8 兼容 ...")
    test_py38(fails)

    if fails:
        print("FAILED:")
        for f in fails:
            print("  - " + f)
        return 1
    print("M13 (票13) SELF-TEST: ALL PASSED "
          "(筒段契约 / 框×桁条自然格 / 开口不连续 / 全链 / 血缘+报告 / 拒绝非法参数 / py3.8)")
    return 0


if __name__ == "__main__":
    sys.exit(main())

# -*- coding: utf-8 -*-
"""
report_template.py -- 生成标准 .docx 报告模板（票 07）

用户之后可替换 `data/templates/report_standard.docx` 为自有规范格式，
只要保留同样的 `{{占位符}}` 即可。

占位符清单见 PLACEHOLDERS。
兼容 Python 3.8。
"""

import os

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None


# 全部占位符（模板生成器与渲染器共用，便于一致性测试）
PLACEHOLDERS = [
    "study_id",
    "generated_at",
    "template_id",
    "summary.max_fi",
    "summary.min_ms",
    "summary.all_pass",
    "summary.n_elems",
    "summary.n_detail_rows",
    "summary.n_bays_checked",
    "summary.n_bays_excluded",
    "check.criteria",
    "check.safety_factor",
    "allowables",
    "param_table",
    "hotspot_table",
    "invalid_bay_table",
    "hotspot_figure",
    "conclusion",
]


def build_standard_template(path):
    """生成一版标准强度校核报告模板（含全部占位符）。"""
    if Document is None:
        raise RuntimeError("python-docx 未安装，无法生成模板")

    doc = Document()

    doc.add_heading("强度校核报告", level=0)

    p = doc.add_paragraph()
    p.add_run("分析任务(Study)：").bold = True
    p.add_run("{{study_id}}")

    p = doc.add_paragraph()
    p.add_run("模板：").bold = True
    p.add_run("{{template_id}}")
    p.add_run("    生成时间：")
    p.add_run("{{generated_at}}")

    # 1 参数与模型
    doc.add_heading("1  参数与模型", level=1)
    doc.add_paragraph("{{param_table}}")

    # 2 校核依据
    doc.add_heading("2  校核依据", level=1)
    p = doc.add_paragraph()
    p.add_run("失效准则：").bold = True
    p.add_run("{{check.criteria}}")
    p = doc.add_paragraph()
    p.add_run("安全系数：").bold = True
    p.add_run("{{check.safety_factor}}")
    p = doc.add_paragraph()
    p.add_run("材料许用值：").bold = True
    p.add_run("{{allowables}}")

    # 3 校核结果汇总
    doc.add_heading("3  校核结果汇总", level=1)
    p = doc.add_paragraph()
    p.add_run("单元数：").bold = True
    p.add_run("{{summary.n_elems}}")
    p.add_run("    明细行数：")
    p.add_run("{{summary.n_detail_rows}}")
    p = doc.add_paragraph()
    p.add_run("校核自然格数：").bold = True
    p.add_run("{{summary.n_bays_checked}}")
    p.add_run("    排除的不连续格数：")
    p.add_run("{{summary.n_bays_excluded}}")
    p = doc.add_paragraph()
    p.add_run("最大失效指数 FI：").bold = True
    p.add_run("{{summary.max_fi}}")
    p = doc.add_paragraph()
    p.add_run("最小安全裕度 MS：").bold = True
    p.add_run("{{summary.min_ms}}")
    p = doc.add_paragraph()
    p.add_run("总体判定：").bold = True
    p.add_run("{{summary.all_pass}}")

    # 4 危险部位
    doc.add_heading("4  危险部位", level=1)
    doc.add_paragraph("{{hotspot_table}}")

    doc.add_heading("4.1  不连续格（已排除出自然格级校核）", level=2)
    doc.add_paragraph("{{invalid_bay_table}}")

    # 5 危险部位图
    doc.add_heading("5  危险部位图", level=1)
    doc.add_paragraph("{{hotspot_figure}}")

    # 6 结论
    doc.add_heading("6  结论与建议", level=1)
    doc.add_paragraph("{{conclusion}}")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("说明：本报告结论文字为 AI 辅助草稿（L1），须经人工审核签发（L2）；"
              "FI/MS 的计算与判定为确定性算法（L0），不含 AI 参与。").italic = True

    out_dir = os.path.dirname(os.path.abspath(path))
    if out_dir and not os.path.isdir(out_dir):
        os.makedirs(out_dir)
    doc.save(path)
    return path


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    repo = os.path.dirname(os.path.dirname(os.path.dirname(here)))
    target = os.path.join(repo, "data", "templates", "report_standard.docx")
    print(build_standard_template(target))

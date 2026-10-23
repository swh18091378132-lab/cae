# -*- coding: utf-8 -*-
"""
test_m20.py -- code-review 复核补丁：spec 中无票认领的三处缺口

这三条要求写在 spec（issue #9）里，但不属于任何一张工单的验收标准，
因此首轮十个实现都没覆盖到。本文件把它们变成可回归的断言：

1. spec 用户故事 18：`为不同部件族指定不同的默认准则与许用值`
   -> 模板 `checks` 声明各族默认口径；`run_study` 以**最低优先级**合并，
      调用方显式传入的 rules 永远优先。
2. spec 用户故事 35：`在报告中看到网格质量警告（若有），以便知道结论的适用边界`
   -> 报告模板有"网格质量与结论适用边界"章节 + `{{mesh_quality_table}}`；
      结论文字在有不达标格时必须提警告。
3. spec 关键契约：`run_study(...) 增加 ... 型号配置引用（可选）`
   -> `RunOptions.model_config_id` 真的被消费：Study 被挂到型号层级，
      型号级血缘能查到它。

另外覆盖一个伴随的正确性问题：标准报告模板有版本标记，
**新增占位符时旧模板必须被重建**，否则新章节在已有安装上永远不出现。

运行 (在 apps/workbench 下): python tests/test_m20.py
"""

import os
import sys
import shutil

HERE = os.path.dirname(os.path.abspath(__file__))
WORKBENCH = os.path.dirname(HERE)
CORE = os.path.join(WORKBENCH, "core")
REPO = os.path.dirname(os.path.dirname(WORKBENCH))
sys.path.insert(0, CORE)
sys.path.insert(0, WORKBENCH)

import db
import study
import template_registry
import model_config
import report_template
import run_study
from run_study import RunOptions
from catia_driver import FakeCatiaDriver
from abaqus_runner import FakeAbaqusRunner


WING_PARAMS = {"span": 5.0, "sweep_deg": 10.0, "c_root": 1.0, "c_tip": 0.6,
               "n_ribs": 3, "n_stringers": 4}


def _tmp(name):
    d = os.path.join(WORKBENCH, ".tmp_test", "m20", name)
    if os.path.isdir(d):
        shutil.rmtree(d)
    os.makedirs(d)
    return d


def _defaults(template_id):
    return dict((p["name"], p.get("default")) for p in
                template_registry.load(template_id).parameters)


def _run(conn, tmp, sid, template_id, params, **extra):
    study.create_study(sid, "M20 " + sid, conn=conn)
    o = RunOptions(
        params_dir=os.path.join(tmp, sid, "params"),
        topo_dir=os.path.join(tmp, sid, "topo"),
        with_fe_model=True, fe_dir=os.path.join(tmp, sid, "fe"),
        subdiv_per_bay=1,
        plies=[{"angle": 0.0, "thickness": 0.1},
               {"angle": 90.0, "thickness": 0.1}],
        abaqus=FakeAbaqusRunner(n_elems=6, n_plies=2),
        with_solve=True, job_name="job1",
        solve_dir=os.path.join(tmp, sid, "solve"),
        with_check=True, check_dir=os.path.join(tmp, sid, "check"),
        with_report=True, report_dir=os.path.join(tmp, sid, "report"))
    for k, v in extra.items():
        o.set(k, v)
    return run_study.run_study(
        sid, params, template_id,
        FakeCatiaDriver(os.path.join(tmp, sid, "cp")), conn, o)


def _doc_text(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# 缺口 1：部件族默认准则与许用值（spec 用户故事 18）
# ---------------------------------------------------------------------------

def test_family_defaults_declared(fails):
    """每个部件族模板都要声明自己的默认校核口径，且各族**可以不同**。"""
    ids = template_registry.list_templates()
    for want in ("wingbox_v1", "tailbox_v1", "fuselage_section_v1"):
        if want not in ids:
            fails.append("模板列表缺少 %s（现有 %r）" % (want, ids))
            continue
        tpl = template_registry.load(want)
        rules = tpl.default_rules()
        if not rules:
            fails.append("%s 未声明 checks（部件族默认口径）" % want)
            continue
        if not rules.get("criteria"):
            fails.append("%s 的 checks 未声明 criteria" % want)
        al = rules.get("allowables") or {}
        for k in ("Xt", "Xc", "Yt", "Yc", "S"):
            if k not in al:
                fails.append("%s 的 checks.allowables 缺 %s" % (want, k))

    # 各族口径不同才叫"贴合各族材料与结构特点"
    wing = template_registry.load("wingbox_v1").default_rules()
    tail = template_registry.load("tailbox_v1").default_rules()
    fuse = template_registry.load("fuselage_section_v1").default_rules()
    combos = [(tuple(wing.get("criteria") or []),),
              (tuple(tail.get("criteria") or []),),
              (tuple(fuse.get("criteria") or []),)]
    if len(set(combos)) == 1 and \
            (wing.get("allowables") == tail.get("allowables")
             == fuse.get("allowables")):
        fails.append("三个部件族的默认准则与许用值完全相同，"
                     "无法体现'各族不同默认口径'（spec 用户故事 18）")

    # 许用值确实不同（至少尾翼与机身 vs 机翼有差异）
    if wing.get("allowables") == tail.get("allowables"):
        fails.append("机翼与尾翼的默认许用值相同，未贴合各族材料特点")


def test_family_defaults_applied(conn, tmp, fails):
    """不传 rules 时用部件族默认口径；传了 rules 时调用方优先。"""
    # 尾翼：用模板默认
    tail = template_registry.load("tailbox_v1")
    tail_defaults = tail.default_rules()
    res = _run(conn, tmp, "S-fam-tail", "tailbox_v1", _defaults("tailbox_v1"))
    summ = res["check"]["summary"]
    if summ.get("criteria") != tail_defaults["criteria"]:
        fails.append("tailbox 未用族默认 criteria: %r != %r"
                     % (summ.get("criteria"), tail_defaults["criteria"]))
    for k, v in tail_defaults["allowables"].items():
        got = (summ.get("allowables") or {}).get(k)
        if got is None or abs(float(got) - float(v)) > 1e-9:
            fails.append("tailbox 许用值 %s 未用族默认: %r != %r" % (k, got, v))
            break
    if abs(float(summ.get("safety_factor")) - float(
            tail_defaults["safety_factor"])) > 1e-9:
        fails.append("tailbox 安全系数未用族默认")

    # 机翼：族默认也生效，且与尾翼**不同**
    res_w = _run(conn, tmp, "S-fam-wing", "wingbox_v1", WING_PARAMS)
    wing_al = res_w["check"]["summary"]["allowables"]
    tail_al = res["check"]["summary"]["allowables"]
    if wing_al.get("Xt") == tail_al.get("Xt"):
        fails.append("机翼与尾翼跑出来的 Xt 相同，族默认未真正分族生效")

    # 调用方显式 rules 优先于族默认
    res_ov = _run(conn, tmp, "S-fam-ov", "tailbox_v1",
                  _defaults("tailbox_v1"),
                  rules={"criteria": ["max_stress"], "safety_factor": 2.0,
                         "allowables": {"Xt": 999.0}})
    s_ov = res_ov["check"]["summary"]
    if s_ov.get("criteria") != ["max_stress"]:
        fails.append("调用方 criteria 未覆盖族默认: %r" % (s_ov.get("criteria"),))
    if abs(float(s_ov.get("safety_factor")) - 2.0) > 1e-9:
        fails.append("调用方 safety_factor 未覆盖族默认")
    if abs(float((s_ov.get("allowables") or {}).get("Xt", 0)) - 999.0) > 1e-9:
        fails.append("调用方 allowables 未覆盖族默认: %r"
                     % (s_ov.get("allowables") or {}).get("Xt"))

    # 规则产物里要能看出用了族默认（可追溯"为什么这套许用值"）
    rule_id = res["artifacts"]["check_rule"]
    row = conn.execute("SELECT params_json FROM artifacts WHERE id=?",
                       (rule_id,)).fetchone()
    if not row or "family_defaults_used" not in (row["params_json"] or ""):
        fails.append("check_rule 产物未记录 family_defaults_used")


# ---------------------------------------------------------------------------
# 缺口 2：报告里的网格质量警告（spec 用户故事 35）
# ---------------------------------------------------------------------------

def test_report_mesh_quality_warning(conn, tmp, fails):
    """有网格质量警告时，报告必须让审核人看到，且结论措辞要受影响。"""
    # 故意把长宽比阈值压到不可能达标，制造必然的告警
    res = _run(conn, tmp, "S-mq", "wingbox_v1", WING_PARAMS,
               mesh_quality={"enabled": True, "auto_refine": False,
                             "aspect_max": 0.5})
    mq = res.get("mesh_quality") or {}
    if not (mq.get("warn_bays") or []):
        fails.append("构造的算例应产生网格质量告警，实际没有")
    rep = res.get("report") or {}
    path = rep.get("path")
    if not path or not os.path.isfile(path):
        fails.append("报告未生成")
        return

    text = _doc_text(path)
    for token in ("网格质量与结论适用边界", "未达标自然格数", "质量警告"):
        if token not in text:
            fails.append("报告缺少网格质量章节要素: %r" % token)
    # 不达标的自然格 ID 要出现在报告里，审核人才能定位
    n_shown = 0
    for bid in (mq.get("warn_bays") or [])[:3]:
        if str(bid) in text:
            n_shown += 1
    if n_shown == 0:
        fails.append("报告未列出任何未达标自然格 ID")
    # 结论文字必须提警告（结论适用边界，不是装饰）
    if "网格质量未达标" not in text:
        fails.append("结论文字未提示网格质量警告（spec 用户故事 35）")

    # 未启用质量门时，报告也要有确定口径（不留空洞/未解析）
    res2 = _run(conn, tmp, "S-mq-off", "wingbox_v1", WING_PARAMS)
    if (res2.get("report") or {}).get("unresolved"):
        fails.append("未启用质量门时报告仍有未解析占位符: %s"
                     % res2["report"]["unresolved"])
    text2 = _doc_text(res2["report"]["path"])
    if "未启用网格质量门" not in text2:
        fails.append("未启用质量门时报告未说明口径")

    # 分析类型（票14）也要出现在报告里，否则审核人不知结论基于什么分析
    if "分析类型" not in text2:
        fails.append("报告缺少分析类型说明")
    res3 = _run(conn, tmp, "S-an", "wingbox_v1", WING_PARAMS,
                analysis="explicit")
    text3 = _doc_text(res3["report"]["path"])
    if "显式动力学" not in text3:
        fails.append("显式分析的报告未写明分析类型")


def test_template_version_rebuild(fails):
    """
    模板版本机制：新增占位符后，**已经存在的旧模板必须被重建**。

    否则 P2 新章节在已有安装上永远不出现——这是真机上最容易被忽略的失效。
    """
    tmp = _tmp("tplver")
    path = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(path)
    if report_template.template_needs_rebuild(path):
        fails.append("刚建好的模板被判为需要重建")
    text = _doc_text(path)
    if report_template.VERSION_MARKER not in text:
        fails.append("模板缺少版本标记 %r" % report_template.VERSION_MARKER)
    # 当前版本必须包含 P2 新增占位符
    for key in ("{{analysis.label}}", "{{mesh_quality.summary}}",
                "{{mesh_quality_table}}"):
        if key not in text:
            fails.append("当前版本模板缺少 P2 占位符 %s" % key)

    # 模拟旧版模板（无版本标记）-> 必须被判为需要重建
    from docx import Document
    old = os.path.join(tmp, "old.docx")
    doc = Document()
    doc.add_paragraph("强度校核报告")
    doc.add_paragraph("{{study_id}}")
    doc.save(old)
    if not report_template.template_needs_rebuild(old):
        fails.append("无版本标记的旧模板未被判为需要重建")
    report_template.ensure_standard_template(old)
    if report_template.template_needs_rebuild(old):
        fails.append("ensure_standard_template 未把旧模板升级到当前版本")
    text_old = _doc_text(old)
    if "{{mesh_quality.summary}}" not in text_old:
        fails.append("升级后的模板仍缺 P2 占位符")

    # 仓库里提交的标准模板必须已是当前版本
    repo_tpl = os.path.join(REPO, "data", "templates", "report_standard.docx")
    if report_template.template_needs_rebuild(repo_tpl):
        fails.append("仓库内的 data/templates/report_standard.docx 是旧版，"
                     "请重新生成（run_study 会重建，但仓库副本也要新）")


# ---------------------------------------------------------------------------
# 缺口 3：run_study 消费 model_config_id（spec 关键契约）
# ---------------------------------------------------------------------------

def test_model_config_reference(conn, tmp, fails):
    """RunOptions.model_config_id 必须真的把 Study 挂到型号层级。"""
    o = RunOptions()
    if not hasattr(o, "model_config_id"):
        fails.append("RunOptions 缺 model_config_id 字段")
        return

    mid = model_config.create_model_config(
        "M20 型号", {"span": 34.0, "fuselage_length": 38.0,
                     "tail_volume": 0.85, "fuselage_diameter": 3.8,
                     "sweep_deg": 25.0}, conn=conn)
    parts = model_config.list_parts(conn, mid)
    wing = None
    for p in parts:
        if p["template_id"] == "wingbox_v1":
            wing = p
            break
    if wing is None:
        fails.append("型号配置未分解出机翼部件")
        return

    res = _run(conn, tmp, "S-mcfg", "wingbox_v1", wing["params"],
               model_config_id=mid)
    mc = res.get("model_config") or {}
    if not mc.get("attached"):
        fails.append("Study 未关联到型号配置: %r（warnings=%r）"
                     % (mc, res.get("warnings")))
    if mc.get("part_key") != wing["part_key"]:
        fails.append("关联到错误部件: %r != %r"
                     % (mc.get("part_key"), wing["part_key"]))

    # 层级里查得到
    lineage = model_config.model_lineage(conn, mid)
    row = [p for p in lineage["parts"]
           if p["template_id"] == "wingbox_v1"]
    if not row or row[0].get("study_id") != "S-mcfg":
        fails.append("型号级血缘未反映关联的 Study: %r"
                     % (row[0] if row else None))

    # 反向查询：从 Study 找到型号
    resolved = model_config.resolve_model_of_study(conn, "S-mcfg")
    got = resolved.get("model_id") if isinstance(resolved, dict) else resolved
    if got != mid:
        fails.append("resolve_model_of_study 未返回型号: %r" % (resolved,))

    # 模板对不上时**不硬塞假部件**，但要显式留痕
    res2 = _run(conn, tmp, "S-mcfg-x", "fuselage_section_v1",
                _defaults("fuselage_section_v1"), model_config_id=mid)
    mc2 = res2.get("model_config") or {}
    if mc2.get("attached") and mc2.get("part_key") == "wing":
        fails.append("模板与部件不对应时错误地挂到了 wing")
    if mc2.get("attached") and mc2.get("part_key") != "fuselage":
        fails.append("机身模板应挂到 fuselage 部件，实际 %r"
                     % mc2.get("part_key"))

    # 不传 model_config_id 时，行为与 P1 完全一致（不影响原有链路）
    res3 = _run(conn, tmp, "S-mcfg-none", "wingbox_v1", WING_PARAMS)
    if "model_config" in res3:
        fails.append("未指定型号配置时不应产生 model_config 字段")

    # 型号 id 不存在时不阻断分析，但必须留痕
    res4 = _run(conn, tmp, "S-mcfg-bad", "wingbox_v1", WING_PARAMS,
                model_config_id="NO-SUCH-MODEL")
    mc4 = res4.get("model_config") or {}
    if mc4.get("attached"):
        fails.append("不存在的型号配置不应关联成功")
    if "report" not in (res4.get("artifacts") or {}):
        fails.append("型号配置关联失败不应阻断报告生成")


def test_analyzing_state_reachable(fails):
    """
    spec 用户故事 26：「看到每个 Study 当前处于流程的哪个环节（分析中/待审核/已签发）」。

    首轮实现里 `analyzing` 只被测试调用过，产品路径永远从 pending_review 起步，
    于是「分析中」是个**到不了的状态**，汇总视图的 n_analyzing 永远是 0。
    修复后：服务层在运行前登记「分析中」，运行结束后转为「待审核」；
    运行失败则撤掉占位，不留幽灵版本。
    """
    import workbench_service as ws
    import approval as ap
    from run_study import RunOptions

    root = _tmp("analysing")
    svc = ws.WorkbenchService(root, mode="fake")
    try:
        sid = svc.start_study("M20 分析中")

        seen = {}

        # 在报告产出**之前**拦截，观察此时的审批状态
        orig = svc._register_report_approval

        def spy(study_id, rep_id, result=None, analysing_version=None):
            st = svc.get_approval(study_id, analysing_version)
            seen["during"] = (st or {}).get("status")
            return orig(study_id, rep_id, result, analysing_version)

        svc._register_report_approval = spy
        svc.run_chain(sid, WING_PARAMS, "wingbox_v1",
                      rules={"criteria": ["tsai_wu"]})
        svc._register_report_approval = orig

        if seen.get("during") != ap.STATUS_ANALYZING:
            fails.append("运行期间 Study 应处于「分析中」，实际 %r"
                         % (seen.get("during"),))

        after = svc.get_approval(sid)
        if (after or {}).get("status") != ap.STATUS_PENDING_REVIEW:
            fails.append("运行结束后应转为「待审核」，实际 %r"
                         % ((after or {}).get("status"),))

        # 运行失败时不留下"分析中"幽灵版本
        sid2 = svc.start_study("M20 失败")
        before = svc.list_approval_versions(sid2)
        try:
            svc.run_chain(sid2, {"span": 9999.0}, "wingbox_v1")
        except ws.ServiceError:
            pass
        else:
            fails.append("越界参数应抛 ServiceError")
        left = [v for v in svc.list_approval_versions(sid2)
                if (v.get("status") if isinstance(v, dict) else None)
                == ap.STATUS_ANALYZING]
        if left:
            fails.append("运行失败后仍留有「分析中」占位: %r" % (left,))
        if before and svc.list_approval_versions(sid2):
            fails.append("运行失败不应留下审批记录")
    finally:
        svc.close()


def test_quality_gate_reachable_from_cli(fails):
    """
    spec "网格生成后自动评估质量指标" —— 必须有一条**产品路径**能启用它，
    否则这道门只有测试会碰，等价于不存在（首轮实现正是如此）。
    """
    import subprocess
    py = sys.executable
    main = os.path.join(WORKBENCH, "main.py")
    proc = subprocess.run([py, main, "run", "wingbox_v1", "--quality-gate"],
                          cwd=WORKBENCH, stdout=subprocess.PIPE,
                          stderr=subprocess.STDOUT, universal_newlines=True)
    if proc.returncode != 0:
        fails.append("CLI --quality-gate 运行失败(rc=%d):\n%s"
                     % (proc.returncode, proc.stdout[-1500:]))
        return
    if "网格质量" not in proc.stdout:
        fails.append("CLI --quality-gate 输出未体现网格质量结果:\n%s"
                     % proc.stdout[-1500:])


def test_py38(fails):
    """本票涉及的模块不得用 3.9+ 语法。"""
    import compat_py38
    targets = [
        os.path.join(CORE, "report.py"),
        os.path.join(CORE, "report_template.py"),
        os.path.join(CORE, "run_study.py"),
        os.path.join(CORE, "template_registry.py"),
        os.path.join(HERE, "test_m20.py"),
    ]
    bad = compat_py38.check_paths(targets)
    if bad:
        fails.append("Python 3.8 兼容性检查未通过:\n%s"
                     % compat_py38.format_violations(bad))


def main():
    fails = []
    tmp = _tmp("db")
    conn = db.connect(os.path.join(tmp, "m20.sqlite"))

    print("1. 部件族默认准则/许用值（声明） ...")
    test_family_defaults_declared(fails)
    print("2. 部件族默认准则/许用值（生效 + 可覆盖） ...")
    test_family_defaults_applied(conn, tmp, fails)
    print("3. 报告中的网格质量警告（spec 用户故事 35） ...")
    test_report_mesh_quality_warning(conn, tmp, fails)
    print("4. 报告模板版本重建 ...")
    test_template_version_rebuild(fails)
    print("5. run_study 消费 model_config_id（spec 关键契约） ...")
    test_model_config_reference(conn, tmp, fails)
    print("6. 「分析中」状态真的可达（spec 用户故事 26） ...")
    test_analyzing_state_reachable(fails)
    print("7. 网格质量门有产品路径可启用 ...")
    test_quality_gate_reachable_from_cli(fails)
    print("8. Python 3.8 兼容 ...")
    test_py38(fails)

    conn.close()

    if fails:
        print("FAILED:")
        for f in fails:
            print("  - " + f)
        return 1
    print("M20 (code-review 补丁) SELF-TEST: ALL PASSED "
          "(族默认口径 / 报告网格质量警告 / 模板版本 / 型号配置引用 / 分析中可达)")
    return 0


if __name__ == "__main__":
    sys.exit(main())

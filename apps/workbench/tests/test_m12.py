# -*- coding: utf-8 -*-
"""
test_m12.py -- 票12 验收：尾翼盒段部件族端到端闭环

验收点（对应 issue #12）：
  1. 模板列表可选到「尾翼盒段」，能读取其参数契约与参数化包络
  2. 给定合法参数集，跑通 建模→拓扑→FE→求解→抽取→校核→报告 全链（fake 适配器）
  3. 自然格按尾翼布置规则生成，数量与规则一致；**操纵面切口**造成的不连续格
     被正确标记且原因正确
  4. 参数越界 / 未知参数按参数契约被拒绝，且**不落半成品产物**
  5. 血缘链完整（参数集→数模→拓扑→网格→inp→odb→抽取→校核结果→报告）
  6. 报告可由该 Study 的校核结果生成，并携带产物指纹（sha256）
  7. Python 3.8 兼容
  8. 机翼盒段零回归（尾翼模板的 cutouts 声明不得泄漏到机翼）

尾翼与机翼同为**翼面类**，但结构特征不同：前后梁位置不固定（由参数给定）、
可能有操纵面切口。本票即验证「换族只要换模板 + 布置规则」这个抽象成立：
本文件除了模板 JSON 外不需要为尾翼写任何专用代码。

运行 (在 apps/workbench 下): python tests/test_m12.py
"""

import os
import sys
import json
import shutil

HERE = os.path.dirname(os.path.abspath(__file__))
WORKBENCH = os.path.dirname(HERE)
CORE = os.path.join(WORKBENCH, "core")
REPO = os.path.dirname(os.path.dirname(WORKBENCH))
sys.path.insert(0, CORE)
sys.path.insert(0, WORKBENCH)

import db
import study
import topology
import template_registry
import run_study
from run_study import RunOptions
from catia_driver import FakeCatiaDriver
from abaqus_runner import FakeAbaqusRunner


TEMPLATE_ID = "tailbox_v1"

# 尾翼盒段合法参数集（= 模板默认值；操纵面切口落在第 3 个肋间、第 2 个长桁间）
PARAMS = {
    "span": 5.0,
    "sweep_deg": 32.0,
    "c_root": 1.8,
    "c_tip": 0.9,
    "n_ribs": 5,
    "n_stringers": 5,
    "spar_front_frac": 0.15,
    "spar_rear_frac": 0.85,
    "cutout_span_frac": 0.6,
    "cutout_chord_frac": 0.375,
}

PLIES = [{"angle": 0.0, "thickness": 0.1},
         {"angle": 90.0, "thickness": 0.1}]

RULES = {"criteria": ["tsai_wu", "hashin"], "safety_factor": 1.0}

# 布置规则推导出的期望值（尾翼参数空间：4 个肋间 × 4 个长桁间 = 16 候选格）
N_RIBS = PARAMS["n_ribs"]
N_STR = PARAMS["n_stringers"]
N_BAYS_TOTAL = (N_RIBS - 1) * (N_STR - 1)

# 操纵面切口位置 -> 落在这一个自然格里（cy 沿翼展分数，cxi 沿前后梁之间分数）
CUTOUT_BAY = "NB-r2-s1"


# ---------------------------------------------------------------------------
# 工具
# ---------------------------------------------------------------------------

def _fresh(name):
    d = os.path.join(WORKBENCH, ".tmp_test", "m12", name)
    if os.path.isdir(d):
        shutil.rmtree(d)
    os.makedirs(d)
    return d


def _tpl():
    return template_registry.load(TEMPLATE_ID)


def _doc_text(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _artifact_count(conn, study_id):
    row = conn.execute("SELECT COUNT(*) c FROM artifacts WHERE study_id=?",
                       (study_id,)).fetchone()
    return int(row["c"])


def _edges_by_kind(conn, study_id):
    """血缘边按 kind 表达（不依赖 id 插入顺序）。"""
    id2kind = {}
    for r in conn.execute("SELECT id, kind FROM artifacts WHERE study_id=?",
                          (study_id,)).fetchall():
        id2kind[r["id"]] = r["kind"]
    edges = set()
    for e in conn.execute("SELECT parent_id, child_id FROM edges").fetchall():
        p, c = id2kind.get(e["parent_id"]), id2kind.get(e["child_id"])
        if p and c:
            edges.add((p, c))
    return edges


def _bay(model, bay_id):
    for b in model.bays:
        if b.id == bay_id:
            return b
    return None


def _inp_node_x_range(inp_path):
    """从 .inp 的 *NODE 段读出物理 X 范围（xi 经 xi_scale 换算后的弦向长度）。"""
    xs = []
    inside = False
    with open(inp_path, "r", encoding="utf-8") as fh:
        for line in fh:
            s = line.strip()
            if s.startswith("*"):
                if inside:
                    break
                inside = s.upper().startswith("*NODE")
                continue
            if not inside or not s:
                continue
            parts = [p.strip() for p in s.split(",")]
            if len(parts) >= 3:
                try:
                    xs.append(float(parts[1]))
                except ValueError:
                    pass
    if not xs:
        return None
    return (min(xs), max(xs))


# ---------------------------------------------------------------------------
# 1) 模板可选到 + 参数契约 / 参数化包络可读
# ---------------------------------------------------------------------------

def test_template_selectable(fails):
    """模板列表能选到尾翼盒段，且参数契约/包络可读。"""
    listed = template_registry.list_templates()
    if TEMPLATE_ID not in listed:
        fails.append("模板列表里没有 %s（现有 %r）" % (TEMPLATE_ID, listed))

    tpl = _tpl()
    if tpl.id != TEMPLATE_ID:
        fails.append("模板 id 不符: %r" % tpl.id)
    if not tpl.name:
        fails.append("模板缺少 name")

    # 参数契约：前后梁位置、翼肋数、翼展、弦长、操纵面切口位置
    contract = tpl.contract()
    required = {
        "span": "翼展",
        "c_root": "根弦长",
        "c_tip": "尖弦长",
        "n_ribs": "翼肋数",
        "n_stringers": "长桁数",
        "spar_front_frac": "前梁位置",
        "spar_rear_frac": "后梁位置",
        "cutout_span_frac": "操纵面切口展向位置",
        "cutout_chord_frac": "操纵面切口弦向位置",
    }
    for name, label in required.items():
        if name not in contract:
            fails.append("参数契约缺少 %s（%s）" % (name, label))
            continue
        spec = contract[name]
        for key in ("min", "max", "default", "type", "map"):
            if key not in spec:
                fails.append("参数 %s 缺少 %s 声明" % (name, key))
        if "min" in spec and "max" in spec and "default" in spec:
            if not (spec["min"] <= spec["default"] <= spec["max"]):
                fails.append("参数 %s 的 default 不在 [min,max] 内: %r"
                             % (name, spec))

    # 参数化包络：必须有跨参数约束（不是逐参数 min/max 就完事）
    constraints = (tpl.envelope or {}).get("constraints") or []
    if len(constraints) < 2:
        fails.append("参数化包络约束过少（%r）" % constraints)
    joined = " ".join(constraints)
    for token in ("c_tip < c_root", "spar_front_frac", "spar_rear_frac"):
        if token not in joined:
            fails.append("参数化包络缺少跨参数约束片段 %r（现有 %r）"
                         % (token, constraints))

    # 拓扑声明：翼面类 + xi 限制在前后梁之间 + 操纵面切口用 holes 表达
    rule = topology.topology_rule(template=tpl)
    if rule.get("mode") != topology.MODE_WING_SURFACE:
        fails.append("尾翼盒段应为翼面类拓扑，实际 %r" % rule.get("mode"))
    rng = (rule.get("ranges") or {}).get("xi") or {}
    if rng.get("from") != "spar_front_frac" or rng.get("to") != "spar_rear_frac":
        fails.append("topology.ranges.xi 应限制在 [spar_front_frac, "
                     "spar_rear_frac]，实际 %r" % (rng,))
    cuts = rule.get("cutouts") or []
    if len(cuts) != 1:
        fails.append("尾翼盒段应声明 1 处操纵面切口，实际 %r" % (cuts,))


# ---------------------------------------------------------------------------
# 2) 自然格按尾翼布置规则生成 + 操纵面切口 -> 不连续格
# ---------------------------------------------------------------------------

def test_layout_and_cutout(fails):
    """自然格数量与布置规则一致；切口造成的不连续格被正确标记且原因正确。"""
    tpl = _tpl()
    params = tpl.validate(PARAMS)
    model = topology.build_topology(TEMPLATE_ID, params, template=tpl)

    # -- 数量与规则：
    #    肋(n_ribs 根，含两端) × 长桁(n_stringers 根，含前后梁) -> (n-1)*(m-1) 格
    if len(model.bays) != N_BAYS_TOTAL:
        fails.append("自然格总数应为 %d，实际 %d（肋 %d / 长桁 %d）"
                     % (N_BAYS_TOTAL, len(model.bays),
                        len(model.ribs), len(model.stringers)))
    if len(model.ribs) != N_RIBS:
        fails.append("翼肋数应为 %d，实际 %d" % (N_RIBS, len(model.ribs)))
    if len(model.stringers) != N_STR:
        fails.append("长桁数应为 %d，实际 %d" % (N_STR, len(model.stringers)))

    # 肋沿翼展等分、覆盖 [0, span]
    ys = [r["y"] for r in model.ribs]
    if abs(ys[0]) > 1e-9 or abs(ys[-1] - params["span"]) > 1e-9:
        fails.append("翼肋未覆盖全翼展: %r" % ys)
    step = params["span"] / float(N_RIBS - 1)
    for k, y in enumerate(ys):
        if abs(y - k * step) > 1e-9:
            fails.append("第 %d 根肋位置不符（等分规则）: %.6f != %.6f"
                         % (k, y, k * step))
            break

    # 长桁只落在前后梁之间（这是尾翼与机翼的结构差别）
    xis = [s["xi"] for s in model.stringers]
    sf, sr = params["spar_front_frac"], params["spar_rear_frac"]
    if abs(xis[0] - sf) > 1e-9 or abs(xis[-1] - sr) > 1e-9:
        fails.append("长桁未限制在前后梁之间: 首 %.6f / 末 %.6f（期望 %.6f / %.6f）"
                     % (xis[0], xis[-1], sf, sr))
    for xi in xis:
        if xi < sf - 1e-9 or xi > sr + 1e-9:
            fails.append("长桁 xi=%.6f 跑出前后梁区间 [%.3f, %.3f]"
                         % (xi, sf, sr))
            break

    # 物理换算：xi 是无量纲分数，乘根弦长才是物理弦向长度
    ext = model.extents()
    if abs(ext["xi_scale"] - params["c_root"]) > 1e-9:
        fails.append("xi_scale 应取根弦长 %.3f，实际 %.3f"
                     % (params["c_root"], ext["xi_scale"]))
    expect_span = (sr - sf) * params["c_root"]
    if abs(ext["xi_span"] - expect_span) > 1e-9:
        fails.append("盒段弦向物理长度应为 (后梁-前梁)*根弦长=%.6f，实际 %.6f"
                     % (expect_span, ext["xi_span"]))

    # -- 操纵面切口 -> 不连续格
    invalid_ids = sorted(b.id for b in model.invalid_bays)
    if invalid_ids != [CUTOUT_BAY]:
        fails.append("操纵面切口应只造成 1 个不连续格 %r，实际 %r"
                     % ([CUTOUT_BAY], invalid_ids))
    bay = _bay(model, CUTOUT_BAY)
    if bay is None:
        fails.append("找不到期望的不连续格 %s" % CUTOUT_BAY)
    else:
        types = [r["type"] for r in bay.reasons]
        if topology.REASON_CUTOUT not in types:
            fails.append("格 %s 的不连续原因里没有 cutout: %r"
                         % (CUTOUT_BAY, bay.reasons))
        if len(types) != 1:
            fails.append("格 %s 应只有 cutout 一个原因，实际 %r"
                         % (CUTOUT_BAY, types))
        # 原因条目要能定位到切口坐标（可追溯，不是只标个类型）
        detail = " ".join(r.get("detail", "") for r in bay.reasons)
        if "cutout at" not in detail:
            fails.append("不连续原因缺少切口位置描述: %r" % detail)
        # 切口确实落在这个格的内部
        hole = model.holes[0]
        if not (bay.y0 < hole["cy"] < bay.y1
                and bay.xi0 < hole["cxi"] < bay.xi1):
            fails.append("切口位置 %r 不在被标记的格 %s 内部 [y %.3f-%.3f, "
                         "xi %.4f-%.4f]" % (hole, CUTOUT_BAY, bay.y0, bay.y1,
                                            bay.xi0, bay.xi1))
        # 切口沿 xi 的分数是"前后梁之间"的分数：映射后应落在前后梁之间
        if not (sf - 1e-9 <= hole["cxi"] <= sr + 1e-9):
            fails.append("切口 cxi=%.4f 未映射进前后梁区间 [%.3f, %.3f]"
                         % (hole["cxi"], sf, sr))

    # -- 反证：去掉模板的 cutouts 声明后应 0 个不连续格
    #    （证明"不连续格由操纵面切口造成"，而不是别的规则副作用）
    rule = topology.topology_rule(template=tpl)
    rule_nocut = dict(rule)
    rule_nocut["cutouts"] = []
    clean = topology.build_topology(TEMPLATE_ID, params, rule=rule_nocut)
    if len(clean.bays) != N_BAYS_TOTAL:
        fails.append("无切口时自然格总数应与规则一致 %d，实际 %d"
                     % (N_BAYS_TOTAL, len(clean.bays)))
    if clean.invalid_bays:
        fails.append("无切口时不应有不连续格，实际 %r"
                     % [b.id for b in clean.invalid_bays])

    # -- 稳定 ID：同参数集两次生成一致且唯一
    again = topology.build_topology(TEMPLATE_ID, params, template=tpl)
    if [b.id for b in model.bays] != [b.id for b in again.bays]:
        fails.append("同参数集两次生成的尾翼自然格 ID 不一致")
    if len(set(b.id for b in model.bays)) != len(model.bays):
        fails.append("尾翼自然格 ID 不唯一")


def test_wing_no_regression(fails):
    """尾翼模板的 cutouts 声明不得泄漏到机翼；机翼盒段零回归。"""
    wing = {"span": 12.0, "sweep_deg": 25.0, "c_root": 3.0, "c_tip": 1.2,
            "n_ribs": 6, "n_stringers": 7}
    m = topology.build_topology("wingbox_v1", wing)
    if len(m.bays) != 30:
        fails.append("机翼零回归: 候选格应为 30，实际 %d" % len(m.bays))
    if m.invalid_bays:
        fails.append("机翼零回归: 干净参数不应有失效格，实际 %r"
                     % [b.id for b in m.invalid_bays])
    if m.holes:
        fails.append("机翼零回归: 机翼模板不应被写入 openings，实际 %r" % m.holes)

    # 机翼的长桁仍是全弦向布置（xi 0..1），没被尾翼的 ranges 影响
    if abs(m.stringers[0]["xi"]) > 1e-9 or abs(m.stringers[-1]["xi"] - 1.0) > 1e-9:
        fails.append("机翼零回归: 长桁未被限制在全弦向 [0,1]，实际 %.6f/%.6f"
                     % (m.stringers[0]["xi"], m.stringers[-1]["xi"]))


# ---------------------------------------------------------------------------
# 3) 端到端全链 + 血缘 + 报告
# ---------------------------------------------------------------------------

def _run_chain(conn, tmp, tag="chain", params=None):
    """用 RunOptions 跑尾翼盒段全链（fake 适配器）。"""
    sid = "S-M12-" + tag
    study.create_study(sid, "M12 尾翼 " + tag, conn=conn)
    opts = RunOptions(
        params_dir=os.path.join(tmp, tag, "params"),
        topo_dir=os.path.join(tmp, tag, "topo"),
        with_topology=True,
        with_fe_model=True,
        fe_dir=os.path.join(tmp, tag, "fe"),
        subdiv_per_bay=1,
        plies=PLIES,
        abaqus=FakeAbaqusRunner(n_elems=N_BAYS_TOTAL, n_plies=2),
        with_solve=True,
        job_name="job1",
        solve_dir=os.path.join(tmp, tag, "solve"),
        with_check=True,
        check_dir=os.path.join(tmp, tag, "check"),
        rules=RULES,
        with_report=True,
        report_dir=os.path.join(tmp, tag, "report"))
    res = run_study.run_study(
        sid, dict(params or PARAMS), TEMPLATE_ID,
        FakeCatiaDriver(os.path.join(tmp, tag, "catparts")), conn, opts)
    return sid, res


def test_full_chain(conn, tmp, fails):
    """合法参数集跑通全链，产物齐全且内容自洽。"""
    sid, res = _run_chain(conn, tmp)

    arts = res.get("artifacts") or {}
    expect_kinds = ("parameter_set", "catpart", "topology", "mesh", "inp",
                    "odb", "extraction", "check_rule", "check_result", "report")
    for kind in expect_kinds:
        if kind not in arts:
            fails.append("全链缺少产物 %s" % kind)
    if res.get("template") != TEMPLATE_ID:
        fails.append("结果的 template 字段不符: %r" % res.get("template"))

    # 参数集按契约规范化后落入产物
    if res.get("parameters") != _tpl().validate(PARAMS):
        fails.append("结果里的参数集未经契约规范化: %r" % (res.get("parameters"),))

    # 拓扑汇总
    topo = res.get("topology") or {}
    if topo.get("n_total") != N_BAYS_TOTAL:
        fails.append("拓扑汇总 n_total 应为 %d，实际 %r"
                     % (N_BAYS_TOTAL, topo.get("n_total")))
    if topo.get("n_invalid") != 1 or topo.get("invalid_ids") != [CUTOUT_BAY]:
        fails.append("拓扑汇总里的不连续格不符: %r" % (topo,))

    # 网格：每格 subdiv=1 -> 每格 1 个单元，且无跨界单元（ADR-0010）
    fe = res.get("fe") or {}
    if fe.get("n_elems") != N_BAYS_TOTAL:
        fails.append("网格单元数应为 %d（每格 1 个），实际 %r"
                     % (N_BAYS_TOTAL, fe.get("n_elems")))
    mesh = res["fe_model"].mesh
    bay_ids = set(b.id for b in res["topology_model"].bays)
    for e in mesh.elems:
        if e["bay_id"] not in bay_ids:
            fails.append("单元 %d 归属未知自然格 %r" % (e["id"], e["bay_id"]))
            break

    # .inp：物理弦向坐标只覆盖前后梁之间那一段（(sr-sf)*c_root）
    inp_path = fe.get("inp_path")
    if not inp_path or not os.path.isfile(inp_path):
        fails.append(".inp 未生成: %r" % inp_path)
    else:
        ext_x = _inp_node_x_range(inp_path)
        expect = (PARAMS["spar_rear_frac"] - PARAMS["spar_front_frac"]) * \
            PARAMS["c_root"]
        if ext_x is None:
            fails.append(".inp 里读不到 *NODE 坐标")
        elif abs((ext_x[1] - ext_x[0]) - expect) > 1e-6:
            fails.append(".inp 弦向跨度应为 %.6f m（前后梁之间），实际 %.6f"
                         % (expect, ext_x[1] - ext_x[0]))

        with open(inp_path, "r", encoding="utf-8") as fh:
            txt = fh.read()
        if ("*ELSET, ELSET=%s" % CUTOUT_BAY) not in txt:
            fails.append(".inp 未按自然格分组（缺 %s 的 ELSET）" % CUTOUT_BAY)

    # 求解 + 抽取
    solve = res.get("solve") or {}
    if solve.get("status") != "COMPLETED":
        fails.append("求解状态应为 COMPLETED，实际 %r" % solve.get("status"))
    extr = res.get("extraction") or {}
    if not extr.get("csv_path") or not os.path.isfile(extr["csv_path"]):
        fails.append("抽取 CSV 未生成: %r" % extr.get("csv_path"))
    manifest = extr.get("manifest") or {}
    if manifest.get("n_rows", 0) <= 0:
        fails.append("抽取 manifest 的 n_rows 应为正数: %r" % manifest)

    # 强度校核：不连续格被排除出自然格级聚合，单元级照常
    chk = res.get("check") or {}
    summ = chk.get("summary") or {}
    if summ.get("n_bays_excluded") != 1:
        fails.append("校核应排除 1 个不连续格，实际 %r"
                     % summ.get("n_bays_excluded"))
    if summ.get("n_bays_checked") != N_BAYS_TOTAL - 1:
        fails.append("校核的有效自然格数应为 %d，实际 %r"
                     % (N_BAYS_TOTAL - 1, summ.get("n_bays_checked")))
    if summ.get("n_detail_rows", 0) <= 0:
        fails.append("校核明细为空")
    if not chk.get("hotspots"):
        fails.append("校核未产出危险部位")
    if _resolve_ids(chk) != [CUTOUT_BAY]:
        fails.append("校核结果里的不连续格清单不符: %r" % _resolve_ids(chk))
    # 明细 CSV 与汇总 JSON 落盘
    for key in ("detail_csv", "summary_json"):
        p = chk.get(key)
        if not p or not os.path.isfile(p):
            fails.append("校核输出 %s 未落盘: %r" % (key, p))

    # 报告
    rep = res.get("report") or {}
    rpath = rep.get("path")
    if not rpath or not os.path.isfile(rpath):
        fails.append("报告未生成: %r" % rpath)
    elif os.path.getsize(rpath) < 5000:
        fails.append("报告过小（%d 字节）" % os.path.getsize(rpath))

    return sid, res


def _resolve_ids(chk):
    ids = chk.get("invalid_bays") or []
    return sorted(ids)


def test_lineage(conn, tmp, sid, res, fails):
    """血缘链完整：参数集→数模→拓扑→网格→inp→odb→抽取→校核结果→报告。"""
    arts = res.get("artifacts") or {}
    edges = _edges_by_kind(conn, sid)

    chain = [("parameter_set", "catpart"),
             ("catpart", "topology"),
             ("topology", "mesh"),
             ("mesh", "inp"),
             ("inp", "odb"),
             ("odb", "extraction"),
             ("extraction", "check_result"),
             ("check_rule", "check_result"),
             ("check_result", "report")]
    for parent, child in chain:
        if parent not in arts or child not in arts:
            fails.append("血缘链缺少端点 %s / %s" % (parent, child))
        elif (parent, child) not in edges:
            fails.append("血缘链缺少边 %s -> %s" % (parent, child))

    # 每个产物都可从报告一路追溯到参数集（参数集是唯一的根）
    nodes = conn.execute(
        "SELECT id, kind, version, checksum, content_uri FROM artifacts "
        "WHERE study_id=? ORDER BY id", (sid,)).fetchall()
    if len(nodes) != len(arts):
        fails.append("血缘节点数 %d 与产物数 %d 不一致" % (len(nodes), len(arts)))
    for n in nodes:
        if not n["checksum"]:
            fails.append("产物 %s 缺少指纹(checksum)" % n["kind"])
    # 版本单调：本次运行的每个 kind 都是 v1
    for n in nodes:
        if n["version"] != 1:
            fails.append("首次运行产物 %s 版本应为 1，实际 %d"
                         % (n["kind"], n["version"]))

    # 血缘 DAG 可从任意节点追溯到参数集
    id2kind = {}
    parents = {}
    for n in nodes:
        id2kind[n["id"]] = n["kind"]
        parents[n["id"]] = []
    for e in conn.execute("SELECT parent_id, child_id FROM edges").fetchall():
        if e["child_id"] in parents:
            parents[e["child_id"]].append(e["parent_id"])

    start = arts["report"]
    seen = set()
    stack = [start]
    while stack:
        cur = stack.pop()
        for p in parents.get(cur, []):
            if p not in seen:
                seen.add(p)
                stack.append(p)
    reached = set(id2kind[i] for i in seen)
    for kind in ("parameter_set", "catpart", "topology", "mesh", "inp",
                 "odb", "extraction", "check_result"):
        if kind not in reached:
            fails.append("从报告无法追溯到 %s" % kind)
    if arts["parameter_set"] not in seen:
        fails.append("从报告无法追溯到参数集（血缘断链）")


def test_report_fingerprint(conn, tmp, sid, res, fails):
    """报告可由该 Study 的校核结果生成，并携带产物指纹。"""
    rep = res.get("report") or {}
    rpath = rep.get("path")
    if not rpath or not os.path.isfile(rpath):
        fails.append("报告不存在，无法校验内容")
        return

    arts = res.get("artifacts") or {}
    # 报告产物本身记录了来源信息
    row = conn.execute(
        "SELECT params_json, version FROM artifacts WHERE id=?",
        (arts["report"],)).fetchone()
    if row is None:
        fails.append("报告产物未入库")
    else:
        payload = json.loads(row["params_json"])
        if payload.get("template") != "report_standard.docx":
            fails.append("报告产物未记录所用模板: %r" % payload.get("template"))
        if int(payload.get("n_hotspots") or 0) <= 0:
            fails.append("报告产物 n_hotspots 应为正数: %r" % payload)

    text = _doc_text(rpath)
    if sid not in text:
        fails.append("报告里没有该 Study 的 id（%s）" % sid)
    if TEMPLATE_ID not in text:
        fails.append("报告里没有记录部件族模板 id（%s）" % TEMPLATE_ID)
    if "强度校核报告" not in text:
        fails.append("报告缺少标题")

    # 产物指纹清单：参数集与校核结果（报告的直接上游）的 sha256 前 12 位必须在报告中
    for kind in ("parameter_set", "check_result", "inp", "extraction"):
        r = conn.execute(
            "SELECT checksum FROM artifacts WHERE id=?", (arts[kind],)).fetchone()
        if r is None:
            fails.append("取不到 %s 的指纹" % kind)
            continue
        prefix = r["checksum"][:12]
        if prefix not in text:
            fails.append("报告的可追溯性附录里没有 %s 的指纹 %s"
                         % (kind, prefix))

    # 报告里写入了尾翼特有参数（证明报告数据来自该 Study 的参数集，不是硬编码）
    if "spar_front_frac" not in text or "spar_rear_frac" not in text:
        fails.append("报告的参数表里没有尾翼的前后梁位置参数")
    if CUTOUT_BAY not in text:
        fails.append("报告的不连续格表里没有操纵面切口造成的 %s" % CUTOUT_BAY)

    # 报告无未解析占位符
    if rep.get("unresolved"):
        fails.append("报告仍有未解析占位符: %r" % rep.get("unresolved"))


# ---------------------------------------------------------------------------
# 4) 参数越界 / 未知参数被拒且不落半成品产物
# ---------------------------------------------------------------------------

def test_rejects(conn, tmp, fails):
    """参数越界/未知参数按契约被拒，且不产生任何产物（不落半成品）。"""
    tpl = _tpl()

    bad_cases = [
        ("单参数越界", dict(PARAMS, span=9999.0)),
        ("单参数低于下限", dict(PARAMS, n_ribs=1)),
        ("未知参数", dict(PARAMS, wing_root_chord=3.0)),
        ("包络违反 c_tip>=c_root", dict(PARAMS, c_tip=2.0)),
        ("包络违反 前后梁顺序", dict(PARAMS, spar_front_frac=0.9)),
        ("包络违反 盒段宽度不足",
         dict(PARAMS, spar_front_frac=0.4, spar_rear_frac=0.5)),
        # 切口落在翼肋线上（0.6 * (6-1) = 3.0 为整数）-> 不落在任何格内部
        ("切口落在翼肋线上", dict(PARAMS, n_ribs=6)),
    ]

    for label, params in bad_cases:
        try:
            tpl.validate(params)
            fails.append("%s: 未被参数契约/包络拒绝 (%r)" % (label, params))
        except template_registry.ParamError:
            pass

    # 端到端：run_study 被拒后不得留下任何产物（也不得建目录/落盘）
    sid_bad = "S-M12-bad"
    study.create_study(sid_bad, "M12 非法参数", conn=conn)
    badout = os.path.join(tmp, "bad")
    bad_params = dict(PARAMS, span=9999.0)
    try:
        run_study.run_study(
            sid_bad, bad_params, TEMPLATE_ID,
            FakeCatiaDriver(os.path.join(badout, "catparts")), conn,
            RunOptions(params_dir=os.path.join(badout, "params"),
                       topo_dir=os.path.join(badout, "topo"),
                       with_fe_model=True,
                       fe_dir=os.path.join(badout, "fe"),
                       abaqus=FakeAbaqusRunner(n_elems=4, n_plies=2),
                       with_solve=True,
                       solve_dir=os.path.join(badout, "solve"),
                       with_check=True,
                       check_dir=os.path.join(badout, "check"),
                       with_report=True,
                       report_dir=os.path.join(badout, "report")))
        fails.append("参数越界的 Study 未被拒绝（应当抛 ParamError）")
    except template_registry.ParamError:
        pass
    except Exception as exc:
        fails.append("参数越界应以 ParamError 拒绝，实际 %s: %s"
                     % (type(exc).__name__, exc))

    n = _artifact_count(conn, sid_bad)
    if n != 0:
        fails.append("参数越界被拒后仍落了 %d 个半成品产物" % n)
    for sub in ("params", "topo", "fe", "solve", "check", "report"):
        d = os.path.join(badout, sub)
        if os.path.isdir(d) and os.listdir(d):
            fails.append("参数越界被拒后仍写出了产物目录 %s: %r"
                         % (sub, os.listdir(d)[:3]))

    # 未知参数同理
    sid_unk = "S-M12-unknown"
    study.create_study(sid_unk, "M12 未知参数", conn=conn)
    try:
        run_study.run_study(
            sid_unk, dict(PARAMS, lead_edge_sweep=10.0), TEMPLATE_ID,
            FakeCatiaDriver(os.path.join(badout, "catparts2")), conn,
            RunOptions(params_dir=os.path.join(badout, "params")))
        fails.append("未知参数未被拒绝")
    except template_registry.ParamError:
        pass
    if _artifact_count(conn, sid_unk) != 0:
        fails.append("未知参数被拒后仍落了产物")

    # 服务层同样把参数错误显式暴露（不吞掉）
    try:
        from workbench_service import WorkbenchService, ServiceError
        root = _fresh("svc")
        svc = WorkbenchService(root, mode="fake")
        try:
            sid_svc = svc.start_study("M12 服务层拒绝")
            try:
                svc.run_chain(sid_svc, dict(PARAMS, c_tip=5.0), TEMPLATE_ID)
                fails.append("服务层未拒绝包络越界的尾翼参数")
            except ServiceError:
                pass
        finally:
            svc.close()
    except ImportError as exc:
        fails.append("服务层导入失败: %s" % exc)


# ---------------------------------------------------------------------------
# 5) Python 3.8 兼容
# ---------------------------------------------------------------------------

def test_py38(fails):
    """本票涉及的模块不得用 3.9+ 语法/标准库。"""
    import compat_py38

    targets = [
        os.path.join(CORE, "topology.py"),
        os.path.join(CORE, "template_registry.py"),
        os.path.join(CORE, "run_study.py"),
        os.path.join(CORE, "compat_py38.py"),
        os.path.join(HERE, "test_m12.py"),
    ]
    bad = compat_py38.check_paths(targets)
    if bad:
        fails.append("Python 3.8 兼容性检查未通过:\n%s"
                     % compat_py38.format_violations(bad))

    # 模板 JSON 必须是可在 3.8 上解析的标准 JSON（不是 JSON5 / 带注释）
    path = os.path.join(REPO, "data", "templates", TEMPLATE_ID + ".json")
    try:
        with open(path, "r", encoding="utf-8") as fh:
            json.load(fh)
    except Exception as exc:
        fails.append("模板 JSON 无法解析: %s" % exc)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main():
    fails = []
    tmp = _fresh("run")
    conn = db.connect(os.path.join(tmp, "m12.db"))

    print("1. 模板可选到 + 参数契约/包络可读 ...")
    test_template_selectable(fails)
    print("2. 尾翼布置规则 + 操纵面切口不连续格 ...")
    test_layout_and_cutout(fails)
    print("3. 机翼盒段零回归 ...")
    test_wing_no_regression(fails)
    print("4. 端到端全链（建模→拓扑→FE→求解→抽取→校核→报告） ...")
    sid, res = test_full_chain(conn, tmp, fails)
    print("5. 血缘链完整 ...")
    test_lineage(conn, tmp, sid, res, fails)
    print("6. 报告携带产物指纹 ...")
    test_report_fingerprint(conn, tmp, sid, res, fails)
    print("7. 参数越界/未知参数被拒且不落半成品 ...")
    test_rejects(conn, tmp, fails)
    print("8. Python 3.8 兼容 ...")
    test_py38(fails)

    conn.close()

    if fails:
        print("FAILED:")
        for f in fails:
            print("  - " + f)
        return 1
    print("M12 (票12) SELF-TEST: ALL PASSED "
          "(%d 自然格 / 1 操纵面切口不连续格 / 全链 fake / 血缘+报告指纹 / 契约拒绝)"
          % N_BAYS_TOTAL)
    return 0


if __name__ == "__main__":
    sys.exit(main())

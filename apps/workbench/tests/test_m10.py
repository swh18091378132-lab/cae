# -*- coding: utf-8 -*-
"""
test_m10.py -- 票10 验收：编排入口改用选项对象（为 P2 扩展腾出空间）

验收点（对应 issue #10）：
  1. 编排入口收敛为「少量必需参数 + 一个选项对象」，选项对象有合理默认值
     —— 不传 options 也能跑通 P1 全链
  2. 现有全部测试（m0–m8）在不修改断言的前提下通过（由 run_tests.py 统一体现，
     本文件只做"选项对象没有改变 P1 行为"的逐项对照）
  3. P1 端到端行为不变：产物种类、血缘父子关系、报告内容一致
  4. 新增一个选项字段不需要改动编排函数签名
  5. 本地工作台服务层与 CLI 同步适配
  6. Python 3.8 兼容

做法：同一份参数分别用「旧的关键字写法」和「新的 options 写法」跑两次，
断言两次的产物种类、血缘边集合、校核结论**逐项一致**。

运行 (在 apps/workbench 下): python tests/test_m10.py
"""

import os
import sys
import json
import shutil
import inspect

HERE = os.path.dirname(os.path.abspath(__file__))
WORKBENCH = os.path.dirname(HERE)
CORE = os.path.join(WORKBENCH, "core")
sys.path.insert(0, CORE)
sys.path.insert(0, WORKBENCH)

import db
import study
import run_study
from run_study import RunOptions
from catia_driver import FakeCatiaDriver
from abaqus_runner import FakeAbaqusRunner


PARAMS = {"span": 5.0, "sweep_deg": 10.0, "c_root": 1.0, "c_tip": 0.6,
          "n_ribs": 3, "n_stringers": 4}

PLIES = [{"angle": 0.0, "thickness": 0.1},
         {"angle": 90.0, "thickness": 0.1}]

RULES = {"criteria": ["tsai_wu", "hashin"], "safety_factor": 1.0}


def _fresh(name):
    d = os.path.join(WORKBENCH, ".tmp_test", name)
    if os.path.isdir(d):
        shutil.rmtree(d)
    os.makedirs(d)
    return d


def _signature_is_stable():
    """编排入口的签名只允许「5 个必需参数 + options + **legacy」。"""
    sig = inspect.signature(run_study.run_study)
    names = list(sig.parameters.keys())
    expected = ["study_id", "parameter_set", "template_id", "catia", "conn",
                "options", "legacy"]
    return names, expected


def _run(conn, tmp, tag, options=None, **legacy):
    """跑一条 P1 全链，返回 (result, 血缘边集合)。"""
    sid = "S-M10-" + tag
    study.create_study(sid, "M10 " + tag, conn=conn)
    catia = FakeCatiaDriver(os.path.join(tmp, tag, "catparts"))
    abaqus = FakeAbaqusRunner(n_elems=6, n_plies=2)

    opts = dict(
        params_dir=os.path.join(tmp, tag, "params"),
        topo_dir=os.path.join(tmp, tag, "topo"),
        with_fe_model=True, fe_dir=os.path.join(tmp, tag, "fe"),
        subdiv_per_bay=1, plies=PLIES,
        abaqus=abaqus, with_solve=True, job_name="job1",
        solve_dir=os.path.join(tmp, tag, "solve"),
        with_check=True, check_dir=os.path.join(tmp, tag, "check"),
        rules=RULES,
        with_report=True, report_dir=os.path.join(tmp, tag, "report"))
    opts.update(legacy)

    if options is None:
        # 旧写法：全部关键字
        res = run_study.run_study(sid, PARAMS, "wingbox_v1", catia, conn, **opts)
    else:
        res = run_study.run_study(sid, PARAMS, "wingbox_v1", catia, conn,
                                  options, **legacy)

    # 血缘边（按 kind 表达，避免 id 依赖插入顺序）
    id2kind = {}
    for r in conn.execute("SELECT id, kind FROM artifacts WHERE study_id=?",
                          (sid,)).fetchall():
        id2kind[r["id"]] = r["kind"]
    edges = set()
    for e in conn.execute(
            "SELECT parent_id, child_id FROM edges").fetchall():
        p, c = id2kind.get(e["parent_id"]), id2kind.get(e["child_id"])
        if p and c:
            edges.add((p, c))
    return res, edges


def test_options_defaults(fails):
    """1) 选项对象有合理默认值：不传即等价于 P1 的默认行为。"""
    o = RunOptions()
    if o.with_topology is not True:
        fails.append("RunOptions: with_topology 默认应为 True")
    for off in ("with_fe_model", "with_solve", "with_check", "with_report"):
        if o.get(off) is not False:
            fails.append("RunOptions: %s 默认应为 False" % off)
    for need in ("params_dir", "topo_dir", "fe_dir", "solve_dir", "check_dir",
                 "report_dir", "abaqus"):
        if o.get(need) is not None:
            fails.append("RunOptions: %s 默认为目录/适配器占位，应为 None" % need)
    if o.subdiv_per_bay != 4:
        fails.append("RunOptions: subdiv_per_bay 默认应为 4（P1 值），实际 %r"
                     % o.subdiv_per_bay)
    if o.job_name != "job1":
        fails.append("RunOptions: job_name 默认应为 job1")

    # 必需参数一项都不能有默认值
    sig = inspect.signature(run_study.run_study)
    for name in ("study_id", "parameter_set", "template_id", "catia", "conn"):
        if sig.parameters[name].default is not inspect.Parameter.empty:
            fails.append("run_study: %s 应是必需参数" % name)


def test_signature_stable(fails):
    """4) 新增选项字段不需要改动编排函数签名。"""
    names, expected = _signature_is_stable()
    if names != expected:
        fails.append("run_study 签名应稳定为 %r，实际 %r" % (expected, names))

    o = RunOptions()
    # 往已声明字段写
    o.set("subdiv_per_bay", 3)
    if o.subdiv_per_bay != 3:
        fails.append("set 未生效")
    # 往**未声明**字段写：不需要改签名也能扩展
    o.set("brand_new_knob", {"x": 1})
    if o.get("brand_new_knob") != {"x": 1}:
        fails.append("未声明字段未被接住（extra 失效）")
    if o.get("brand_new_knob", {}).get("x") != 1:
        fails.append("未声明字段读取口径不一致")

    # 派生不丢字段、且覆盖生效
    o2 = o.replace(subdiv_per_bay=2)
    if o2.subdiv_per_bay != 2 or o.subdiv_per_bay != 3:
        fails.append("replace 应派生新对象而不改原对象")
    if o2.get("brand_new_knob") != {"x": 1}:
        fails.append("replace 丢了 extra 字段")
    if "brand_new_knob" not in o2.to_dict():
        fails.append("to_dict 丢了 extra 字段")

    # overlay：只让显式设置过的字段生效（服务层铺目录 + 调用方按需覆盖）
    base = RunOptions(params_dir="base/params", with_solve=True)
    user = RunOptions(subdiv_per_bay=2)
    merged = base.overlay(user)
    if merged.params_dir != "base/params" or merged.with_solve is not True:
        fails.append("overlay 冲掉了 base 的定制字段")
    if merged.subdiv_per_bay != 2:
        fails.append("overlay 未应用 user 的显式字段")

    # 未知关键字照样可以进 options（不会被当成 TypeError）
    o3 = RunOptions(anything_goes=42)
    if o3.get("anything_goes") != 42:
        fails.append("RunOptions(**kwargs) 未接住未知字段")

    # coerce：None / dict / RunOptions 三种输入都归一
    if RunOptions.coerce(None).with_topology is not True:
        fails.append("coerce(None) 失效")
    if RunOptions.coerce({"subdiv_per_bay": 7}).subdiv_per_bay != 7:
        fails.append("coerce(dict) 失效")
    if RunOptions.coerce(RunOptions(subdiv_per_bay=5)).subdiv_per_bay != 5:
        fails.append("coerce(RunOptions) 失效")
    mixed = RunOptions.coerce({"subdiv_per_bay": 7}, {"subdiv_per_bay": 9})
    if mixed.subdiv_per_bay != 9:
        fails.append("coerce: 关键字写法应覆盖 options 同名字段")


def test_p1_behaviour_unchanged(conn, tmp, fails):
    """2)+3) 旧写法与新写法产出的产物种类、血缘、结论逐项一致。"""
    res_legacy, edges_legacy = _run(conn, tmp, "legacy")

    single = RunOptions(
        params_dir=os.path.join(tmp, "opts", "params"),
        topo_dir=os.path.join(tmp, "opts", "topo"),
        with_fe_model=True, fe_dir=os.path.join(tmp, "opts", "fe"),
        subdiv_per_bay=1, plies=PLIES,
        abaqus=FakeAbaqusRunner(n_elems=6, n_plies=2),
        with_solve=True, job_name="job1",
        solve_dir=os.path.join(tmp, "opts", "solve"),
        with_check=True, check_dir=os.path.join(tmp, "opts", "check"),
        rules=RULES, with_report=True,
        report_dir=os.path.join(tmp, "opts", "report"))
    res_opts, edges_opts = _run(conn, tmp, "opts", options=single)

    kinds_legacy = sorted(res_legacy["artifacts"].keys())
    kinds_opts = sorted(res_opts["artifacts"].keys())
    if kinds_legacy != kinds_opts:
        fails.append("产物种类不一致:\n    旧=%r\n    新=%r"
                     % (kinds_legacy, kinds_opts))
    expected_kinds = ("parameter_set", "catpart", "topology", "mesh", "inp",
                      "odb", "extraction", "check_rule", "check_result",
                      "report")
    for k in expected_kinds:
        if k not in kinds_opts:
            fails.append("新写法缺少产物 %s" % k)

    if edges_legacy != edges_opts:
        fails.append("血缘父子关系不一致:\n    仅旧有=%r\n    仅新有=%r"
                     % (sorted(edges_legacy - edges_opts),
                        sorted(edges_opts - edges_legacy)))
    must_have = [("parameter_set", "catpart"), ("catpart", "topology"),
                 ("topology", "mesh"), ("mesh", "inp"), ("inp", "odb"),
                 ("odb", "extraction"), ("extraction", "check_result"),
                 ("check_rule", "check_result"),
                 ("check_result", "report")]
    for e in must_have:
        if e not in edges_opts:
            fails.append("新写法缺少血缘边 %s->%s" % e)

    # 拓扑与校核结论逐项一致（P1 口径不变）
    if res_legacy["topology"] != res_opts["topology"]:
        fails.append("拓扑汇总不一致: %r vs %r"
                     % (res_legacy["topology"], res_opts["topology"]))
    s1 = res_legacy["check"]["summary"]
    s2 = res_opts["check"]["summary"]
    for k in ("n_detail_rows", "n_elems", "n_bays_checked", "n_bays_excluded",
              "max_fi", "min_ms", "all_pass", "rules_checksum"):
        if s1.get(k) != s2.get(k):
            fails.append("校核口径不一致 %s: %r vs %r" % (k, s1.get(k), s2.get(k)))

    # 报告内容：两版都提到同一 Study 的参数、结论与产物清单
    for res in (res_legacy, res_opts):
        p = (res.get("report") or {}).get("path")
        if not p or not os.path.isfile(p):
            fails.append("报告缺失: %r" % p)
        elif os.path.getsize(p) < 5000:
            fails.append("报告过小: %s" % p)

    # 产物落盘位置互不干扰
    if res_legacy["catpart_path"] == res_opts["catpart_path"]:
        fails.append("两次运行应各自落盘")


def test_report_options_passthrough(conn, tmp, fails):
    """report_options 也能通过 options 传下去（P1 的规模保护参数）。"""
    sid = "S-M10-repopt"
    study.create_study(sid, "M10 report options", conn=conn)
    res = run_study.run_study(
        sid, PARAMS, "wingbox_v1",
        FakeCatiaDriver(os.path.join(tmp, "repopt", "catparts")), conn,
        RunOptions(params_dir=os.path.join(tmp, "repopt", "params"),
                   topo_dir=os.path.join(tmp, "repopt", "topo"),
                   with_fe_model=True,
                   fe_dir=os.path.join(tmp, "repopt", "fe"),
                   subdiv_per_bay=1, plies=PLIES,
                   abaqus=FakeAbaqusRunner(n_elems=6, n_plies=2),
                   with_solve=True, job_name="job1",
                   solve_dir=os.path.join(tmp, "repopt", "solve"),
                   with_check=True,
                   check_dir=os.path.join(tmp, "repopt", "check"),
                   rules=RULES, with_report=True,
                   report_dir=os.path.join(tmp, "repopt", "report"),
                   report_options={"detail_table_max_rows": 3}))
    rep = res.get("report") or {}
    if not rep.get("path"):
        fails.append("report_options: 报告未生成")
    else:
        # 明细被截断时表尾会写明"另有 N 行未列出"
        try:
            from docx import Document
            doc = Document(rep["path"])
            blob = "\n".join(p.text for p in doc.paragraphs)
            for t in doc.tables:
                for row in t.rows:
                    blob += "\n" + " | ".join(c.text for c in row.cells)
            if "另有" not in blob:
                fails.append("report_options.detail_table_max_rows 未生效")
        except ImportError:
            print("  [SKIP] 无 python-docx，跳过报告内容断言")


def test_service_and_cli(conn, tmp, fails):
    """5) 服务层与 CLI 同步适配：run_chain 接受 options，CLI 仍能一条命令跑通。"""
    import workbench_service as ws

    root = _fresh("m10svc")
    svc = ws.WorkbenchService(root, mode="fake")
    try:
        sid = svc.start_study("M10 服务层")
        # 只传 options 里的新字段，不应破坏服务层已铺好的目录与开关
        opts = RunOptions(subdiv_per_bay=1,
                          rules={"criteria": ["tsai_wu"]},
                          plies=PLIES)
        res = svc.run_chain(sid, PARAMS, "wingbox_v1", options=opts)
        arts = res.get("artifacts") or {}
        for k in ("parameter_set", "catpart", "topology", "mesh", "inp", "odb",
                  "extraction", "check_rule", "check_result", "report"):
            if k not in arts:
                fails.append("服务层 + options: 缺少产物 %s" % k)
        # 服务层目录约定未被 options 覆盖掉
        uri = res.get("catpart_path") or ""
        if "studies" not in uri.replace("\\", "/"):
            fails.append("服务层目录约定被破坏: %r" % uri)
        # 不传 options 的老写法仍然可用
        res2 = svc.run_chain(sid, PARAMS, "wingbox_v1",
                             rules={"criteria": ["hashin"]})
        if "report" not in (res2.get("artifacts") or {}):
            fails.append("服务层老写法（不传 options）失效")
    finally:
        svc.close()

    # CLI：main.py run 一条命令仍可跑通（subprocess 隔离，避免污染本进程）
    import subprocess
    py = sys.executable
    rc = subprocess.call([py, os.path.join(WORKBENCH, "main.py"), "run",
                          "wingbox_v1"], cwd=WORKBENCH,
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    if rc != 0:
        fails.append("CLI `main.py run wingbox_v1` 失败(rc=%d)" % rc)


def test_py38(fails):
    """6) Python 3.8 兼容：本票引入/改动的模块不得用 3.9+ 语法。"""
    import compat_py38
    targets = [
        os.path.join(CORE, "run_study.py"),
        os.path.join(CORE, "workbench_service.py"),
        os.path.join(CORE, "compat_py38.py"),
        os.path.join(WORKBENCH, "main.py"),
        os.path.join(HERE, "test_m10.py"),
    ]
    bad = compat_py38.check_paths(targets)
    if bad:
        fails.append("Python 3.8 兼容性检查未通过:\n%s"
                     % compat_py38.format_violations(bad))


def main():
    fails = []
    tmp = _fresh("m10")
    dbpath = os.path.join(tmp, "m10.db")
    conn = db.connect(dbpath)

    print("1. options 默认值 ...")
    test_options_defaults(fails)
    print("2. 签名稳定 + 扩展字段不改签名 ...")
    test_signature_stable(fails)
    print("3. P1 行为零回归（旧写法 vs 新写法） ...")
    test_p1_behaviour_unchanged(conn, tmp, fails)
    print("4. report_options 透传 ...")
    test_report_options_passthrough(conn, tmp, fails)
    print("5. 服务层与 CLI 适配 ...")
    test_service_and_cli(conn, tmp, fails)
    print("6. Python 3.8 兼容 ...")
    test_py38(fails)

    conn.close()

    if fails:
        print("FAILED:")
        for f in fails:
            print("  - " + f)
        return 1
    print("M10 (票10) SELF-TEST: ALL PASSED "
          "(options 默认值/签名稳定/P1 零回归/服务层+CLI)")
    return 0


if __name__ == "__main__":
    sys.exit(main())

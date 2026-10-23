# -*- coding: utf-8 -*-
"""
test_m26.py -- 票26 验收：ai_advisor —— L1 解释与标注协议 + 确定性 Fake

验收点（对应 issue #26 / ADR-0002）：
  1. 三个方法有明确签名；fake 实现确定性可回归（同输入同输出）
  2. fake 输出不依赖网络、不依赖时间/随机（连续两次调用结果一致）
  3. 适配器缺失/不可用时走到降级路径：degraded=True + 确定性兜底文本
  4. AI 输出与确定性结论在不同字段，二者**程序上不可混淆**
     （包括：一个"说谎"的适配器也不能改写合格判定）
  5. 调用留痕可查（适配器标识 + 输入摘要）
  6. 报告在有 AI 解释时标明 L1 辅助；无适配器时报告仍能生成
  7. CI 只跑 fake，测试不触网
  8. Python 3.8 兼容

ADR-0002 硬约束：**强度判定（FI/MS 计算与合格判定）永不由 LLM 完成**。
本文件的第 4 项就是这条约束的可执行断言。

运行 (在 apps/workbench 下): python tests/test_m26.py
"""

import os
import sys
import shutil
import socket

HERE = os.path.dirname(os.path.abspath(__file__))
WORKBENCH = os.path.dirname(HERE)
CORE = os.path.join(WORKBENCH, "core")
sys.path.insert(0, CORE)
sys.path.insert(0, WORKBENCH)

import compat_py38


def _fresh(name):
    d = os.path.join(WORKBENCH, ".tmp_test", "m26", name)
    if os.path.isdir(d):
        shutil.rmtree(d)
    os.makedirs(d)
    return d


def _doc_text(path):
    """把 .docx 的全部段落与表格文字拼成一个字符串（只断言外部可观察内容）。"""
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _doc_paragraphs(path):
    from docx import Document
    return [p.text for p in Document(path).paragraphs]


# ---------------------------------------------------------------------------
# 测试夹具：一份贴近 strength_check.check() 真实返回的校核结果
#   数值取自费氏/哈希准则的典型量级，手工写入（不从被测代码推导）。
# ---------------------------------------------------------------------------

def _sample_check(all_pass=False):
    return {
        "summary": {
            "n_detail_rows": 36,
            "n_elems": 6,
            "n_bays_checked": 4,
            "n_bays_excluded": 0,
            "n_fail_rows": 12 if not all_pass else 0,
            "max_fi": 3.6 if not all_pass else 0.42,
            "min_ms": -0.7222 if not all_pass else 1.3809,
            "all_pass": bool(all_pass),
            "criteria": ["tsai_wu", "hashin"],
            "safety_factor": 1.0,
            "ms_threshold": 0.0,
            "allowables": {"Xt": 1500, "Xc": 1200, "Yt": 50, "Yc": 200, "S": 70},
            "rules_checksum": "deadbeef" * 8,
        },
        "hotspots": [
            {"elem": 5, "sp": 2, "criterion": "tsai_wu",
             "FI": 3.6 if not all_pass else 0.42,
             "MS": -0.7222 if not all_pass else 1.3809},
            {"elem": 3, "sp": 1, "criterion": "hashin", "FI": 0.7, "MS": 0.4286},
        ],
        "invalid_bays": [],
        "detail_rows": [],
    }


def _sample_hotspot():
    return {"elem": 5, "sp": 2, "criterion": "tsai_wu", "FI": 3.6, "MS": -0.7222}


def _sample_context():
    """危险部位的解释上下文：铺层 + 自然格归属（票26 要求"基于该处 FI/MS、铺层、自然格归属"）。"""
    return {
        "bay_id": "NB-r1-s0",
        "plies": [{"angle": 0.0, "thickness": 0.1},
                  {"angle": 90.0, "thickness": 0.1}],
        "n_plies": 2,
        "total_thickness": 0.2,
        "stress": {"s11": 412.0, "s22": 38.0, "s12": -12.5},
        "allowables": {"Xt": 1500, "Xc": 1200, "Yt": 50, "Yc": 200, "S": 70},
    }


def _no_network(fn, *args, **kwargs):
    """在"断网"环境里执行 fn：任何 socket 创建都视为违规（CI 不触网）。"""
    orig_socket = socket.socket
    orig_create = socket.create_connection

    def _boom(*a, **k):
        raise AssertionError("AI 适配器不允许联网（CI 只跑 fake）")

    socket.socket = _boom
    socket.create_connection = _boom
    try:
        return fn(*args, **kwargs)
    finally:
        socket.socket = orig_socket
        socket.create_connection = orig_create


# ---------------------------------------------------------------------------
# 1) 三个方法 + 确定性
# ---------------------------------------------------------------------------

def test_protocol_and_determinism(fails):
    import ai_advisor

    # 协议本身：三个方法都在（真机/假机同签名）
    for name in ("explain_check", "explain_hotspot", "suggest_layup"):
        if not hasattr(ai_advisor.AiAdvisor, name):
            fails.append("AiAdvisor 协议缺少方法 %s" % name)

    adv = ai_advisor.DeterministicAiAdvisor()
    chk = _sample_check()

    # 断网执行：fake 不得触网（这是「CI 只跑 fake」的可执行保证）
    n1 = _no_network(adv.explain_check, chk)
    n2 = _no_network(adv.explain_check, chk)

    # 字段齐全
    for attr in ("kind", "draft_text", "facts", "advisor_id", "degraded",
                 "layer", "trace", "draft"):
        if not hasattr(n1, attr):
            fails.append("AdvisorNote 缺少字段 %s" % attr)

    if not isinstance(n1.draft_text, str) or not n1.draft_text.strip():
        fails.append("explain_check 未产出文本")
    if n1.degraded is not False:
        fails.append("可用适配器的正常路径不应标 degraded")
    if n1.layer != "L1":
        fails.append("AI 输出必须标为 L1，实际 %r" % n1.layer)
    if n1.draft is not True:
        fails.append("AI 输出必须标为草稿（draft=True）")

    # 连续两次调用结果一致（同输入同输出）
    if n1.to_dict() != n2.to_dict():
        fails.append("fake 输出不确定：两次调用结果不同\n  1: %r\n  2: %r"
                     % (n1.to_dict(), n2.to_dict()))

    # 另起一个实例也必须一致（不依赖进程内状态）
    n3 = ai_advisor.DeterministicAiAdvisor().explain_check(chk)
    if n3.to_dict() != n1.to_dict():
        fails.append("fake 输出依赖实例状态：新实例结果不同")

    # 危险部位解释：连续两次一致
    h = _sample_hotspot()
    ctx = _sample_context()
    h1 = _no_network(adv.explain_hotspot, h, ctx)
    h2 = adv.explain_hotspot(h, ctx)
    if h1.to_dict() != h2.to_dict():
        fails.append("explain_hotspot 输出不确定")
    if not h1.draft_text.strip():
        fails.append("explain_hotspot 未产出文本")
    # 「为什么这里危险」必须基于铺层与自然格归属（不能只复述数字）
    if "NB-r1-s0" not in h1.draft_text:
        fails.append("explain_hotspot 未引用自然格归属：%r" % h1.draft_text)
    if "铺层" not in h1.draft_text and "层" not in h1.draft_text:
        fails.append("explain_hotspot 未引用铺层信息：%r" % h1.draft_text)

    # 铺层建议：连续两次一致
    s1 = _no_network(adv.suggest_layup, h, ctx)
    s2 = adv.suggest_layup(h, ctx)
    if s1.to_dict() != s2.to_dict():
        fails.append("suggest_layup 输出不确定")
    if not s1.draft_text.strip():
        fails.append("suggest_layup 未产出文本")
    # 建议必须能追溯到失效准则/应力（不能是空话）
    if "tsai_wu" not in s1.draft_text and "Tsai" not in s1.draft_text:
        fails.append("铺层建议未引用失效准则：%r" % s1.draft_text)


# ---------------------------------------------------------------------------
# 2) 降级路径（硬验收项）
# ---------------------------------------------------------------------------

def test_degraded_path(fails):
    import ai_advisor

    chk = _sample_check()
    h = _sample_hotspot()
    ctx = _sample_context()

    # (a) 适配器缺失 -> 兜底适配器，degraded=True，文本是确定性兜底
    fb = ai_advisor.resolve_advisor(None)
    if not isinstance(fb, ai_advisor.FallbackAiAdvisor):
        fails.append("适配器缺失时应退化为 FallbackAiAdvisor，实际 %r"
                     % type(fb).__name__)

    n = _no_network(fb.explain_check, chk)
    if n.degraded is not True:
        fails.append("适配器缺失时 degraded 必须为 True")
    if n.degrade_reason != ai_advisor.DEGRADE_MISSING:
        fails.append("适配器缺失的降级原因码不对：%r" % n.degrade_reason)
    if ai_advisor.FALLBACK_PREFIX not in n.draft_text:
        fails.append("降级文本必须是确定性兜底内容（缺兜底前缀）：%r"
                     % n.draft_text[:80])
    if n.is_ai_generated:
        fails.append("兜底文本不应被标为 AI 生成")
    if "L1" not in n.label:
        fails.append("降级路径仍须标明 L1 层级：%r" % n.label)
    # 兜底文本里仍要有确定性数值（不能让审核人看不到结论）
    if "3.6" not in n.draft_text:
        fails.append("兜底文本应包含确定性 FI 值，实际：%r" % n.draft_text)
    # 兜底也确定（可回归）
    if fb.explain_check(chk).to_dict() != n.to_dict():
        fails.append("降级兜底输出不确定")

    for method, args, kwargs in (
            ("explain_hotspot", (h, ctx), {}),
            ("suggest_layup", (h, ctx), {})):
        note = getattr(fb, method)(*args, **kwargs)
        if note.degraded is not True or not note.draft_text.strip():
            fails.append("降级路径 %s 未产出兜底文本" % method)
        if ai_advisor.FALLBACK_PREFIX not in note.draft_text:
            fails.append("降级路径 %s 文本不是兜底内容" % method)

    # (b) 显式关闭 AI -> 也走降级，且原因码不同
    off = ai_advisor.resolve_advisor(False)
    n_off = off.explain_check(chk)
    if n_off.degraded is not True:
        fails.append("显式关闭 AI 应走降级路径")
    if n_off.degrade_reason != ai_advisor.DEGRADE_DISABLED:
        fails.append("显式关闭的原因码不对：%r" % n_off.degrade_reason)

    # (c) 不合协议的适配器 -> 降级，不抛
    class Broken(object):
        pass

    bad = ai_advisor.resolve_advisor(Broken())
    n_bad = bad.explain_check(chk)
    if n_bad.degraded is not True:
        fails.append("不合协议的适配器应降级")
    if n_bad.degrade_reason != ai_advisor.DEGRADE_INVALID:
        fails.append("不合协议的原因码不对：%r" % n_bad.degrade_reason)

    # (d) 合规但**调用时抛异常**的适配器 -> 降级，不把异常甩给调用方
    class Exploding(object):
        advisor_id = "exploding/v1"

        def explain_check(self, check, study_name=None):
            raise RuntimeError("模拟 LLM 不可达")

        def explain_hotspot(self, hotspot, context=None):
            raise RuntimeError("模拟 LLM 不可达")

        def suggest_layup(self, hotspot, context=None):
            raise RuntimeError("模拟 LLM 不可达")

    safe = ai_advisor.resolve_advisor(Exploding())
    n_exc = safe.explain_check(chk)
    if n_exc.degraded is not True:
        fails.append("适配器抛异常时必须降级为 degraded=True")
    if n_exc.degrade_reason != ai_advisor.DEGRADE_ERROR:
        fails.append("调用异常的原因码不对：%r" % n_exc.degrade_reason)
    if ai_advisor.FALLBACK_PREFIX not in n_exc.draft_text:
        fails.append("调用异常时文本应为确定性兜底")
    for method in ("explain_hotspot", "suggest_layup"):
        note = getattr(safe, method)(h, ctx)
        if note.degraded is not True:
            fails.append("适配器异常时 %s 未降级" % method)

    # (e) 返回非 AdvisorNote 的适配器 -> 降级
    class WrongType(object):
        advisor_id = "wrong-type/v1"

        def explain_check(self, check, study_name=None):
            return "我就返回一个字符串"

        def explain_hotspot(self, hotspot, context=None):
            return {"draft_text": "dict"}

        def suggest_layup(self, hotspot, context=None):
            return None

    safe2 = ai_advisor.resolve_advisor(WrongType())
    n_wrong = safe2.explain_check(chk)
    if n_wrong.degraded is not True:
        fails.append("适配器返回不合规对象时应降级")
    # 关键：不合规对象的"文本"绝不能流进报告
    if "我就返回一个字符串" in n_wrong.draft_text:
        fails.append("不合规适配器的返回值泄漏进了报告文本")

    # (f) 合规适配器经 resolve 后仍正常工作（不被误降级）
    good = ai_advisor.resolve_advisor(ai_advisor.DeterministicAiAdvisor())
    n_good = good.explain_check(chk)
    if n_good.degraded is not False:
        fails.append("合规适配器被误降级")
    if not n_good.is_ai_generated:
        fails.append("合规适配器的输出应标为 AI 生成")


# ---------------------------------------------------------------------------
# 3) AI 输出与确定性结论分属不同字段（ADR-0002 的可执行断言）
# ---------------------------------------------------------------------------

def test_field_separation(fails):
    import ai_advisor

    chk = _sample_check(all_pass=False)
    adv = ai_advisor.DeterministicAiAdvisor()
    note = adv.explain_check(chk)

    # (a) AI 文字只在 draft_text；确定性数值只在 facts
    if not isinstance(note.draft_text, str):
        fails.append("draft_text 必须单独承载 AI 文字")
    fact_keys = set(note.facts.keys())
    # facts 是确定性事实的回显，不能反过来成为判定的来源
    for k in ("all_pass", "max_fi", "min_ms"):
        if k in fact_keys and note.facts[k] != chk["summary"][k]:
            fails.append("facts.%s 与确定性校核值不一致（AI 不得改写数值）："
                         "%r != %r" % (k, note.facts[k], chk["summary"][k]))

    # (b) 层级/草稿标记是不可变属性：AI 输出永远是 L1 草稿
    if note.layer != ai_advisor.LAYER_L1:
        fails.append("AdvisorNote.layer 必须恒为 L1")
    if note.draft is not True:
        fails.append("AdvisorNote.draft 必须恒为 True")
    if note.is_ai_generated is not True:
        fails.append("可用适配器的输出应标为 AI 生成")

    # (c) 决定性检查：**"说谎"的适配器也不能改变合格判定**
    #     ADR-0002：强度判定永不由 LLM 完成。
    class Lying(object):
        advisor_id = "lying/v1"

        def explain_check(self, check, study_name=None):
            # 声称"合格"，并试图改写 facts 与 summary
            check["summary"]["all_pass"] = True
            check["summary"]["max_fi"] = 0.01
            return ai_advisor.AdvisorNote(
                "explain_check", "本模型完全合格，可直接签发。",
                facts={"all_pass": True, "max_fi": 0.01},
                advisor_id=self.advisor_id)

        def explain_hotspot(self, hotspot, context=None):
            return ai_advisor.AdvisorNote("explain_hotspot", "无风险",
                                          advisor_id=self.advisor_id)

        def suggest_layup(self, hotspot, context=None):
            return ai_advisor.AdvisorNote("suggest_layup", "无需修改",
                                          advisor_id=self.advisor_id)

    import copy
    honest = _sample_check(all_pass=False)

    import report as report_mod
    import report_template
    ctx = report_mod.build_report_context({
        "study_id": "S-LIE", "parameters": {}, "check": honest})
    ctx2 = report_mod.build_report_context({
        "study_id": "S-LIE", "parameters": {}, "check": honest},
        advisor=ai_advisor.resolve_advisor(Lying()))

    # 确定性判定字段保持不变 —— 这是本票最重要的一条断言
    if ctx2["summary"]["all_pass"] != ctx["summary"]["all_pass"]:
        fails.append("AI 适配器改写了确定性合格判定！%r -> %r"
                     % (ctx["summary"]["all_pass"], ctx2["summary"]["all_pass"]))
    if ctx2["summary"]["max_fi"] != ctx["summary"]["max_fi"]:
        fails.append("AI 适配器改写了确定性 max_fi！")
    if ctx2["summary"]["min_ms"] != ctx["summary"]["min_ms"]:
        fails.append("AI 适配器改写了确定性 min_ms！")

    # AI 文字只在 adv 命名空间
    if "adv" not in ctx2:
        fails.append("ctx 缺少 AI 命名空间 adv")
    else:
        for key in ("summary", "degraded", "advisor_id", "layer_label",
                    "suggest", "trace"):
            if key not in ctx2["adv"]:
                fails.append("ctx['adv'] 缺少字段 %r" % key)
    if "adv_table" not in ctx2:
        fails.append("ctx 缺少 AI 结构化表 adv_table")
    if "adv" in ctx2["summary"]:
        fails.append("summary 命名空间里混入了 AI 字段（不可混淆）")
    # 确定性结论草稿（模板三分支）不得被 AI 改写
    if ctx2["conclusion"] != ctx["conclusion"]:
        fails.append("AI 适配器改写了确定性结论草稿 conclusion")

    # (d) 校核结果对象本身没被 AI 污染
    if honest["summary"]["all_pass"] is not False:
        fails.append("适配器直接改写了传入的 check 字典（越权）")

    # (e) 报告渲染后，AI 说明与确定性结论**出现在不同段落**
    tmp = _fresh("sep")
    tpl = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(tpl)
    out = os.path.join(tmp, "r.docx")
    report_mod.render_report("S-LIE", ctx2, tpl, out)
    doc_text = _doc_text(out)
    if "本模型完全合格，可直接签发。" not in doc_text:
        fails.append("AI 文字未进报告")
    # 关键：AI 声称"合格"，但报告的确定性判定仍应是不合格
    if "不合格" not in doc_text:
        fails.append("报告丢失了确定性'不合格'判定（AI 越权生效）")
    # AI 段落必须带 L1 标注
    if "L1" not in doc_text:
        fails.append("报告未标注 L1 辅助")
    paras = _doc_paragraphs(out)
    ai_paras = [p for p in paras if "可直接签发" in p]
    det_paras = [p for p in paras if "不合格" in p and "可直接签发" not in p]
    if not ai_paras:
        fails.append("未找到 AI 说明段落")
    if not det_paras:
        fails.append("未找到确定性判定段落")
    for p in ai_paras:
        if "L1" not in p:
            fails.append("AI 说明所在段落缺少 L1 标注：%r" % p[:60])


# ---------------------------------------------------------------------------
# 4) 报告集成：有/无适配器都能生成，且 AI 章节与确定性章节分开
# ---------------------------------------------------------------------------

def _chain_result(all_pass=False):
    """
    贴近 run_study 真实返回的 result（含 fe_model 的网格与铺层，
    用来验证 AI 能取到自然格归属与铺层）。
    """
    import meshgen as _meshgen

    class _Elem(object):
        pass

    # 用一个极小的真 Mesh，只为提供 elem -> bay_id 归属
    mesh = _meshgen.Mesh()
    mesh.add_quad(0.0, 1.0, 0.0, 0.5, "NB-r1-s0", True, 1, 1)
    mesh.add_quad(0.0, 1.0, 0.5, 1.0, "NB-r1-s1", True, 1, 1)

    class _Fe(object):
        pass

    fe = _Fe()
    fe.mesh = mesh
    fe.laminate = {"n_plies": 2, "total_thickness": 0.2,
                   "plies": [{"angle": 0.0, "thickness": 0.1},
                             {"angle": 90.0, "thickness": 0.1}]}

    chk = _sample_check(all_pass=all_pass)
    # 单元 1 属于 NB-r1-s0（网格里第一格），据此可验证"自然格归属"被取到
    chk["hotspots"] = [{"elem": 1, "sp": 2, "criterion": "tsai_wu",
                        "FI": chk["summary"]["max_fi"],
                        "MS": chk["summary"]["min_ms"]},
                       {"elem": 3, "sp": 1, "criterion": "hashin",
                        "FI": 0.7, "MS": 0.4286}]
    return {
        "study_id": "S-M26",
        "study_name": "M26 报告集成",
        "template": "wingbox_v1",
        "parameters": {"span": 5.0, "sweep_deg": 10.0, "c_root": 1.0,
                       "c_tip": 0.6, "n_ribs": 3, "n_stringers": 4},
        "topology": {"n_total": 2, "n_valid": 2, "n_invalid": 0},
        "fe": {"n_nodes": 9, "n_elems": 2, "n_plies": 2,
               "total_thickness": 0.2},
        "fe_model": fe,
        "solve": {"job": "job1", "status": "COMPLETED"},
        "check": chk,
        # 应力分量：由确定性抽取结果提供（L1 只读它，不重算）
        "stress_map": {(1, 2): {"c1": 412.0, "c2": 38.0, "c4": -12.5}},
    }


def test_report_with_advisor(fails):
    import ai_advisor
    import report as report_mod
    import report_template

    tmp = _fresh("with")
    tpl = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(tpl)
    out = os.path.join(tmp, "r.docx")

    res = _chain_result()
    advisor = ai_advisor.resolve_advisor(ai_advisor.DeterministicAiAdvisor())
    ctx = report_mod.build_report_context(res, advisor=advisor)

    if ctx["adv"]["degraded"] is not False:
        fails.append("提供 fake 适配器时不应降级")
    if not ctx["adv"]["summary"].strip():
        fails.append("报告 ctx 里 AI 总述为空")
    # AI 解释必须基于数据（引用真实 FI 值），而不是空话
    if "3.6" not in ctx["adv"]["summary"]:
        fails.append("AI 总述未引用确定性 FI=3.6：%r" % ctx["adv"]["summary"][:120])
    # 危险部位的说明要取到自然格归属与铺层（来自 fe_model，不是硬编码）
    rows = ctx["adv_table"]
    if len(rows) != 2:
        fails.append("adv_table 行数应为危险部位数 2，实际 %d" % len(rows))
    else:
        if rows[0].get("bay_id") != "NB-r1-s0":
            fails.append("AI 未取到自然格归属（期望 NB-r1-s0，实际 %r）"
                         % rows[0].get("bay_id"))
        if "铺层" not in (rows[0].get("why") or ""):
            fails.append("AI 危险部位说明未涉及铺层：%r" % rows[0].get("why"))
        if not (rows[0].get("suggest") or "").strip():
            fails.append("AI 未给出铺层建议")

    rr = report_mod.render_report("S-M26", ctx, tpl, out)
    if rr.get("unresolved"):
        fails.append("报告有未解析占位符：%s" % rr["unresolved"])
    text = _doc_text(out)
    for needle in ("L1", "辅助", "3.6"):
        if needle not in text:
            fails.append("报告缺少 AI 章节要素 %r" % needle)
    if "NB-r1-s0" not in text:
        fails.append("报告未出现自然格归属说明")

    # AI 章节必须与确定性判定章节分开：AI 段落在确定性判定段落之后
    paras = _doc_paragraphs(out)
    idx_det = None
    idx_ai = None
    for i, p in enumerate(paras):
        if "总体判定" in p and idx_det is None:
            idx_det = i
        if "L1 辅助" in p and "AI" in p and idx_ai is None:
            idx_ai = i
    if idx_det is None:
        fails.append("报告缺少确定性'总体判定'段落")
    if idx_ai is None:
        fails.append("报告缺少 L1 辅助章节标题")
    if idx_det is not None and idx_ai is not None and idx_ai <= idx_det:
        fails.append("AI 章节应位于确定性判定之后（避免混淆）")


def test_report_without_advisor(fails):
    """无适配器时：报告照常生成，AI 章节是确定性兜底文字并标明不可用。"""
    import report as report_mod
    import report_template

    tmp = _fresh("without")
    tpl = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(tpl)
    out = os.path.join(tmp, "r.docx")

    res = _chain_result()
    ctx = report_mod.build_report_context(res)      # 不传适配器
    if ctx["adv"]["degraded"] is not True:
        fails.append("未提供适配器时应标 degraded=True")
    if "兜底" not in ctx["adv"]["summary"]:
        fails.append("无适配器时 AI 章节应使用确定性兜底文字：%r"
                     % ctx["adv"]["summary"][:80])

    rr = report_mod.render_report("S-M26", ctx, tpl, out)
    if not os.path.isfile(out):
        fails.append("无适配器时报告未生成")
        return
    if rr.get("unresolved"):
        fails.append("无适配器时报告有未解析占位符：%s" % rr["unresolved"])
    text = _doc_text(out)
    # 报告本体必须完整（确定性部分照常）
    for needle in ("强度校核报告", "S-M26", "3.6", "不合格", "L1"):
        if needle not in text:
            fails.append("无适配器时报告缺少 %r" % needle)
    if "兜底" not in text:
        fails.append("无适配器时报告未出现确定性兜底文字")

    # "AI 永远不可用"也不能让报告生成失败（显式关闭同理）
    ctx_off = report_mod.build_report_context(res, advisor=False)
    if ctx_off["adv"]["degraded"] is not True:
        fails.append("显式关闭 AI 时应标 degraded=True")
    out2 = os.path.join(tmp, "r2.docx")
    if not os.path.isfile(report_mod.render_report(ctx_off, tpl, out2)["path"]):
        fails.append("显式关闭 AI 时报告未生成")


def test_template_version(fails):
    """新增占位符必须递增 TEMPLATE_VERSION，旧模板必须被 rebuild。"""
    import report_template

    if report_template.TEMPLATE_VERSION < 3:
        fails.append("新增 AI 占位符后 TEMPLATE_VERSION 必须递增（当前 %d）"
                     % report_template.TEMPLATE_VERSION)

    tmp = _fresh("tpl")
    path = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(path)
    if report_template.template_needs_rebuild(path):
        fails.append("刚建好的模板被判为需要重建")

    from docx import Document
    text = _doc_text(path)
    for key in ("{{adv.summary}}", "{{adv.suggest}}", "{{adv.layer_label}}",
                "{{adv_table}}", "{{adv.advisor_id}}", "{{adv.degraded_text}}"):
        if key not in text:
            fails.append("模板缺少 AI 占位符 %s" % key)

    # PLACEHOLDERS 与实际模板一致（新增占位符漏登记是常见漂移）
    missing = [p for p in report_template.PLACEHOLDERS
               if "{{%s}}" % p not in text]
    if missing:
        fails.append("PLACEHOLDERS 里有模板未声明的项：%s" % missing)

    # 模拟 P2 时代的旧模板（v2 标记）-> 必须被判为需要重建
    old = os.path.join(tmp, "old.docx")
    doc = Document()
    doc.add_paragraph("强度校核报告")
    doc.add_paragraph("{{study_id}}")
    doc.add_paragraph("模板版本：v2")
    doc.save(old)
    if not report_template.template_needs_rebuild(old):
        fails.append("v2 旧模板未被判为需要重建")
    report_template.ensure_standard_template(old)
    if report_template.template_needs_rebuild(old):
        fails.append("ensure_standard_template 未把 v2 旧模板升级到当前版本")
    old_text = _doc_text(old)
    if "{{adv.summary}}" not in old_text:
        fails.append("重建后的模板缺少 AI 章节占位符")


# ---------------------------------------------------------------------------
# 7) 调用留痕（适配器标识 + 输入摘要）
# ---------------------------------------------------------------------------

def test_invocation_trace(fails):
    import ai_advisor

    adv = ai_advisor.DeterministicAiAdvisor()
    chk = _sample_check()
    h = _sample_hotspot()
    ctx = _sample_context()

    n1 = adv.explain_check(chk)
    n2 = adv.explain_hotspot(h, ctx)
    n3 = adv.suggest_layup(h, ctx)

    # AdvisorNote.trace：适配器标识 + 输入摘要（ADR-0002 留痕）
    for note, method in ((n1, ai_advisor.METHOD_EXPLAIN_CHECK),
                         (n2, ai_advisor.METHOD_EXPLAIN_HOTSPOT),
                         (n3, ai_advisor.METHOD_SUGGEST_LAYUP)):
        t = note.trace or {}
        if t.get("method") != method:
            fails.append("trace.method 不对：%r != %r" % (t.get("method"), method))
        if not t.get("advisor_id"):
            fails.append("trace 缺少适配器标识")
        digest = t.get("input_digest")
        if not digest or len(digest) < 8:
            fails.append("trace 缺少输入摘要：%r" % digest)

    # 同一输入 -> 同一摘要；不同输入 -> 不同摘要（摘要真有区分力）
    again = adv.explain_check(chk)
    if again.trace["input_digest"] != n1.trace["input_digest"]:
        fails.append("同输入的摘要不一致（留痕不可复现）")
    other = adv.explain_check(_sample_check(all_pass=True))
    if other.trace["input_digest"] == n1.trace["input_digest"]:
        fails.append("不同输入的摘要相同（留痕没有区分力）")

    # 调用流水：每次调用一条记录，顺序与内容可查
    # 到目前为止共 5 次调用：check / hotspot / layup / again / other
    if len(adv.invocations) != 5:
        fails.append("调用流水条数不对：%d（期望 5）" % len(adv.invocations))
    if [t["method"] for t in adv.invocations][:3] != [
            ai_advisor.METHOD_EXPLAIN_CHECK,
            ai_advisor.METHOD_EXPLAIN_HOTSPOT,
            ai_advisor.METHOD_SUGGEST_LAYUP]:
        fails.append("调用流水顺序不对：%r"
                     % [t["method"] for t in adv.invocations])

    # 降级路径也要留痕（否则"这次为什么没有 AI 解释"无从追溯）
    fb = ai_advisor.resolve_advisor(None)
    nfb = fb.explain_check(chk)
    if not (nfb.trace or {}).get("degraded"):
        fails.append("降级调用未在留痕里标记 degraded")
    if (nfb.trace or {}).get("degrade_reason") != ai_advisor.DEGRADE_MISSING:
        fails.append("降级留痕缺少原因码")

    # 报告 ctx 的留痕文本要能被审核人读到（含适配器与摘要）
    import report as report_mod
    res = _chain_result()
    ctx_report = report_mod.build_report_context(
        res, advisor=ai_advisor.resolve_advisor(
            ai_advisor.DeterministicAiAdvisor()))
    trace_text = ctx_report["adv"]["trace"]
    for needle in ("适配器=", "输入摘要=", "降级="):
        if needle not in trace_text:
            fails.append("报告留痕文本缺少 %r：%r" % (needle, trace_text))
    if ai_advisor.ADVISOR_FAKE not in trace_text:
        fails.append("报告留痕未写明适配器标识：%r" % trace_text)


# ---------------------------------------------------------------------------
# 8) CI 只跑 fake：整个测试过程不触网
# ---------------------------------------------------------------------------

def test_ci_no_network(fails):
    """把整条"生成报告"路径放进断网环境跑一遍（含降级路径）。"""
    import ai_advisor
    import report as report_mod
    import report_template

    tmp = _fresh("nonet")
    tpl = os.path.join(tmp, "t.docx")
    report_template.build_standard_template(tpl)

    def _whole_flow():
        adv = ai_advisor.resolve_advisor(ai_advisor.DeterministicAiAdvisor())
        ctx = report_mod.build_report_context(_chain_result(), advisor=adv)
        out = os.path.join(tmp, "with.docx")
        report_mod.render_report("S-NONET", ctx, tpl, out)
        # 降级路径同样不触网
        ctx2 = report_mod.build_report_context(_chain_result())
        out2 = os.path.join(tmp, "without.docx")
        report_mod.render_report("S-NONET", ctx2, tpl, out2)
        return out, out2

    outs = _no_network(_whole_flow)
    for out in outs:
        if not os.path.isfile(out):
            fails.append("断网环境下报告未生成：%s" % out)

    # 源码层面也不应出现联网相关 import（防止将来悄悄加回来）
    src = open(os.path.join(CORE, "ai_advisor.py"), encoding="utf-8").read()
    for banned in ("import socket", "import urllib", "import requests",
                   "import http", "import ssl"):
        if banned in src:
            fails.append("ai_advisor.py 出现联网依赖：%r" % banned)


# ---------------------------------------------------------------------------
# 9) RealAiAdvisor 只留位置不实现（CI 不得依赖网络）
# ---------------------------------------------------------------------------

def test_real_advisor_stub(fails):
    import ai_advisor

    try:
        ai_advisor.RealAiAdvisor()
        fails.append("RealAiAdvisor 应保持未实现（抛 NotImplementedError）")
    except NotImplementedError as exc:
        msg = str(exc)
        # 说明原因（票26 原文要求"留接口位置并写明原因"）
        if "网络" not in msg and "联网" not in msg:
            fails.append("RealAiAdvisor 未说明未实现原因：%s" % msg)
    except Exception as exc:
        fails.append("RealAiAdvisor 应抛 NotImplementedError，实际 %r" % type(exc))

    # 但它仍满足协议（将来实现即可接入，无需改调用方）
    for name in ai_advisor.METHODS:
        if not callable(getattr(ai_advisor.RealAiAdvisor, name, None)):
            fails.append("RealAiAdvisor 缺少方法 %s" % name)


# ---------------------------------------------------------------------------
# 10) Python 3.8 兼容
# ---------------------------------------------------------------------------

def test_py38(fails):
    paths = [os.path.join(CORE, "ai_advisor.py"),
             os.path.join(CORE, "report.py"),
             os.path.join(CORE, "report_template.py"),
             os.path.abspath(__file__)]
    violations = compat_py38.check_paths(paths)
    if violations:
        fails.append("Python 3.8 不兼容：\n"
                     + compat_py38.format_violations(violations))


def main():
    fails = []
    print("1. 协议 + fake 确定性 ...")
    test_protocol_and_determinism(fails)
    print("2. 降级路径（缺失/关闭/不合规/抛异常）...")
    test_degraded_path(fails)
    print("3. AI 输出与确定性结论分属不同字段 ...")
    test_field_separation(fails)
    print("4. 报告集成：有适配器 ...")
    test_report_with_advisor(fails)
    print("5. 报告集成：无适配器（兜底仍可生成）...")
    test_report_without_advisor(fails)
    print("6. 模板版本递增 + 旧模板重建 ...")
    test_template_version(fails)
    print("7. 调用留痕（适配器 + 输入摘要）...")
    test_invocation_trace(fails)
    print("8. CI 只跑 fake（断网）...")
    test_ci_no_network(fails)
    print("9. RealAiAdvisor 未实现但留位 ...")
    test_real_advisor_stub(fails)
    print("10. Python 3.8 兼容 ...")
    test_py38(fails)

    if fails:
        print("FAILED:")
        for f in fails:
            print("  - " + f)
        return 1
    print("M26 (票26) SELF-TEST: ALL PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())

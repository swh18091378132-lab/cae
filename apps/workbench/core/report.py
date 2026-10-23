# -*- coding: utf-8 -*-
"""
report.py -- 强度校核报告渲染（票 07）

读 .docx 模板，替换 {{占位符}}，插入表格与图片，输出报告。

设计要点：
  * 文本占位符：{{key}}，支持点号路径 {{summary.max_fi}}
  * 结构化占位符：{{param_table}} / {{hotspot_table}} / {{invalid_bay_table}}
    在所在段落位置插入真实 Word 表格
  * 图片占位符：{{hotspot_figure}} 在所在段落插入图片
  * 未知占位符：**保留原样**并记入 warnings（不静默丢弃）

兼容 Python 3.8。
"""

import os
import re
import datetime

try:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:  # pragma: no cover
    Document = None

from criteria import DEFAULT_ALLOWABLES

PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z0-9_.\-]+)\s*\}\}")

# 渲染期选项（由 render_report 设置），用于规模保护等
_REPORT_OPTIONS = {}

TABLE_PLACEHOLDERS = ("param_table", "detail_table", "hotspot_table",
                      "invalid_bay_table", "artifact_table")
IMAGE_PLACEHOLDERS = ("hotspot_figure",)


# ---------------------------------------------------------------------------
# context 构造
# ---------------------------------------------------------------------------

def build_report_context(result, generated_at=None):
    """
    从 run_study 的返回 dict 构造报告 context。

    result: run_study 返回值（需含 parameters / check / topology 等）

    占位符命名空间（与票 07 约定一致）：
      {{study_id}} {{study_name}} {{template_id}} {{generated_at}}
      {{param.span}} {{param.n_ribs}} ...        逐参数
      {{model.n_elems}} {{model.n_nodes}} ...    网格/铺层/自然格规模
      {{check.criteria}} {{check.safety_factor}} {{check.allowables}} ...
      {{summary.max_fi}} {{summary.all_pass}} ...
      结构化：{{param_table}} {{detail_table}} {{hotspot_table}}
             {{invalid_bay_table}} {{artifact_table}} {{hotspot_figure}}
    """
    chk = result.get("check") or {}
    summ = chk.get("summary") or {}
    params = result.get("parameters") or {}
    topo = result.get("topology") or {}
    fe = result.get("fe") or {}
    solve = result.get("solve") or {}

    allow = summ.get("allowables") or DEFAULT_ALLOWABLES
    allow_txt = ("Xt=%g, Xc=%g, Yt=%g, Yc=%g, S=%g MPa"
                 % (allow.get("Xt", 0), allow.get("Xc", 0),
                    allow.get("Yt", 0), allow.get("Yc", 0), allow.get("S", 0)))

    crit = summ.get("criteria") or []
    all_pass = summ.get("all_pass")
    max_fi = summ.get("max_fi")
    min_ms = summ.get("min_ms")

    # 结论草稿（L1：AI 辅助文字，人审核签发 L2）
    if all_pass is True:
        verdict = "经校核，本模型在给定准则与安全系数下满足强度要求。"
        advice = "建议维持当前设计方案；后续可针对质量指标开展优化迭代。"
    elif all_pass is False:
        verdict = ("经校核，本模型存在不满足强度要求的部位（最大失效指数 FI=%.3g）。"
                   % (max_fi if max_fi is not None else 0.0))
        advice = ("建议对第 4 章所列危险部位进行局部加强（增加铺层/改变铺向/加筋），"
                  "并重新校核验证。")
    else:
        verdict = "本报告未包含有效校核结果。"
        advice = "请检查求解与抽取环节是否成功。"

    conclusion = (verdict + "\n" + advice +
                  "\n（本结论为 AI 辅助草稿 L1，须经强度分析人员审核后签发 L2。）")

    hotspots = chk.get("hotspots") or []
    invalid_bays = chk.get("invalid_bays") or []
    detail_rows = chk.get("detail_rows") or []

    ctx = {
        "study_id": result.get("study_id", ""),
        "study_name": result.get("study_name") or result.get("study_id", ""),
        "template_id": result.get("template", ""),
        "generated_at": generated_at or datetime.datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"),

        # 逐参数命名空间：{{param.span}}
        "param": dict(params),

        # 模型信息命名空间：{{model.n_elems}}
        "model": {
            "n_bays_total": topo.get("n_total", 0),
            "n_bays_valid": topo.get("n_valid", 0),
            "n_bays_invalid": topo.get("n_invalid", 0),
            "n_elems": fe.get("n_elems", 0),
            "n_nodes": fe.get("n_nodes", 0),
            "n_plies": fe.get("n_plies", 0),
            "total_thickness": fe.get("total_thickness", 0),
            "job": solve.get("job", ""),
            "status": solve.get("status", ""),
        },

        "summary": {
            "max_fi": _fmt(max_fi),
            "min_ms": _fmt(min_ms),
            "all_pass": ("合格" if all_pass is True
                         else ("不合格" if all_pass is False else "未知")),
            "n_elems": summ.get("n_elems", 0),
            "n_detail_rows": summ.get("n_detail_rows", 0),
            "n_bays_checked": summ.get("n_bays_checked", 0),
            "n_bays_excluded": summ.get("n_bays_excluded", 0),
            "n_fail_rows": summ.get("n_fail_rows", 0),
        },
        "check": {
            "criteria": "、".join(crit) if crit else "（未指定）",
            "safety_factor": _fmt(summ.get("safety_factor")),
            "ms_threshold": _fmt(summ.get("ms_threshold")),
            "allowables": allow_txt,
            "rules_checksum": summ.get("rules_checksum", ""),
        },

        # 结构化（表格/图片）
        "param_table": params,
        "detail_table": detail_rows,
        "hotspot_table": hotspots,
        "invalid_bay_table": invalid_bays,
        "artifact_table": _artifact_rows(result),
        "hotspot_figure": result.get("hotspot_figure"),
        "conclusion": conclusion,

        # 兼容旧模板的顶层别名（{{allowables}} 等）
        "allowables": allow_txt,
    }
    return ctx


def _artifact_rows(result):
    """把 result['artifacts'] 整理成可追溯的清单（kind/version/checksum）。"""
    arts = result.get("artifacts") or {}
    meta = result.get("artifact_meta") or {}
    rows = []
    for kind in sorted(arts.keys()):
        info = meta.get(kind) or {}
        rows.append({
            "kind": kind,
            "version": info.get("version", ""),
            "checksum": (info.get("checksum") or "")[:12],
            "uri": info.get("uri", ""),
        })
    return rows


def _fmt(v):
    if v is None:
        return "—"
    try:
        return "%.4g" % float(v)
    except (TypeError, ValueError):
        return str(v)


def _lookup(ctx, key):
    """按点号路径取值；缺失返回 None。"""
    cur = ctx
    for part in key.split("."):
        if isinstance(cur, dict) and part in cur:
            cur = cur[part]
        else:
            return None
    return cur


# ---------------------------------------------------------------------------
# 渲染
# ---------------------------------------------------------------------------

def render_report(*args, **kwargs):
    """
    渲染报告。返回 {'path':..., 'warnings':[...], 'replaced': n, 'unresolved':[...]}

    支持两种调用形式（票 07 约定的在前）：
        render_report(study_id, context, template_path, out_path)
        render_report(context, template_path, out_path)      # 兼容旧调用

    可选 report_options（关键字）：
        detail_table_max_rows : 明细表最大行数（0/None = 不限）。真实算例明细可达
                                几十万行，全灌入 docx 会撑爆；超出时保留 FI 最高者
                                并在表尾注明，全量仍见 check_detail.csv。
    """
    global _REPORT_OPTIONS
    opts = kwargs.pop("report_options", None)
    _REPORT_OPTIONS = opts or {}

    if len(args) >= 4:
        study_id, context, template_path, out_path = args[:4]
    elif len(args) == 3:
        study_id = kwargs.pop("study_id", None)
        context, template_path, out_path = args
    else:
        raise TypeError("render_report 需要 (study_id, context, template_path, out_path)"
                        " 或 (context, template_path, out_path)")

    if study_id:
        # 显式 study_id 优先（追溯：报告必须能对上具体 Study 版本）
        context = dict(context)
        context["study_id"] = study_id

    if Document is None:
        raise RuntimeError("python-docx 未安装，无法渲染报告")
    if not os.path.isfile(template_path):
        raise FileNotFoundError("报告模板不存在: %s" % template_path)

    doc = Document(template_path)
    warnings = []
    replaced = 0

    # ---- 1) 段落级：文本占位符 + 表格/图片占位符 ----
    for para in list(doc.paragraphs):
        replaced += _process_paragraph(doc, para, context, warnings)

    # ---- 2) 表格内文本占位符 ----
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replaced += _replace_text_in_para(para, context, warnings)

    out_dir = os.path.dirname(os.path.abspath(out_path))
    if out_dir and not os.path.isdir(out_dir):
        os.makedirs(out_dir)

    doc.save(out_path)
    return {"path": out_path, "warnings": warnings, "replaced": replaced,
            "unresolved": _scan_unresolved(out_path)}


def _process_paragraph(doc, para, context, warnings):
    """处理一个段落：图片占位符 > 表格占位符 > 文本占位符。返回替换数。"""
    text = para.text or ""

    # 图片占位符
    for key in IMAGE_PLACEHOLDERS:
        token = "{{%s}}" % key
        if token in text or ("{{ %s }}" % key) in text:
            img = _lookup(context, key)
            if img and os.path.isfile(img):
                _clear_para(para)
                try:
                    run = para.add_run()
                    run.add_picture(img, width=Cm(15.5))
                except Exception as exc:
                    warnings.append("图片插入失败 %s: %s" % (key, exc))
            else:
                warnings.append("图片占位符 %s 无可用图片（%r）" % (key, img))
                _clear_para(para)
                para.add_run("[图缺失: %s]" % key)
            return 1

    # 表格占位符
    for key in TABLE_PLACEHOLDERS:
        if ("{{%s}}" % key) in text or ("{{ %s }}" % key) in text:
            data = _lookup(context, key)
            _clear_para(para)
            _insert_table_at(doc, para, key, data, warnings)
            return 1

    # 纯文本占位符
    return _replace_text_in_para(para, context, warnings)


def _replace_text_in_para(para, context, warnings):
    """替换段落内的 {{key}}；保留原格式尽量不变（简化为整段重建）。"""
    text = para.text or ""
    if "{{" not in text:
        return 0

    count = [0]

    def sub(m):
        key = m.group(1)
        val = _lookup(context, key)
        if val is None or isinstance(val, (dict, list)):
            # 结构化占位符不该走到这里；未知则保留原样
            warnings.append("未解析占位符: {{%s}}" % key)
            return m.group(0)
        count[0] += 1
        return str(val)

    new_text = PLACEHOLDER_RE.sub(sub, text)
    if new_text != text:
        _set_para_text(para, new_text)
    return count[0]


def _clear_para(para):
    for run in list(para.runs):
        run._element.getparent().remove(run._element)


def _set_para_text(para, text):
    _clear_para(para)
    run = para.add_run(text)
    try:
        run.font.size = Pt(10.5)
    except Exception:
        pass


def _insert_table_at(doc, para, key, data, warnings):
    """在指定段落之后插入表格（简化：追加到文档末尾前，再移动到位置）。"""
    if key == "param_table":
        rows = [("参数", "取值")]
        if isinstance(data, dict):
            for k in sorted(data.keys()):
                rows.append((str(k), str(data[k])))
        else:
            warnings.append("param_table 数据非 dict")
    elif key == "hotspot_table":
        rows = [("序号", "单元", "铺层", "准则", "FI", "MS")]
        for i, h in enumerate(data or [], 1):
            rows.append((str(i), str(h.get("elem")),
                         str(h.get("sp")), str(h.get("criterion")),
                         _fmt(h.get("FI")), _fmt(h.get("MS"))))
    elif key == "detail_table":
        # 校核明细表（单元 × 铺层 × 准则），FI 降序便于审阅
        # 规模保护：真实算例可达几十万行，全灌进 docx 会撑爆。
        # 可通过 report_options.detail_table_max_rows 设上限（0/None = 不限）。
        rows = [("单元", "铺层", "准则", "FI", "MS", "判定")]
        items = sorted(data or [], key=lambda r: -(r.get("FI") or 0.0))
        limit = int((_REPORT_OPTIONS or {}).get("detail_table_max_rows") or 0)
        truncated = 0
        if limit > 0 and len(items) > limit:
            truncated = len(items) - limit
            items = items[:limit]
        for r in items:
            rows.append((str(r.get("elem")), str(r.get("sp")),
                         str(r.get("criterion")),
                         _fmt(r.get("FI")), _fmt(r.get("MS")),
                         "合格" if r.get("pass") else "不合格"))
        if truncated:
            rows.append(("…", "…", "…", "…", "…",
                         "另有 %d 行未列出（全量见 check_detail.csv）" % truncated))
    elif key == "artifact_table":
        # 产物与指纹清单：报告内容可追溯到具体版本
        rows = [("产物", "版本", "指纹(sha256)", "位置")]
        for a in (data or []):
            rows.append((str(a.get("kind")), str(a.get("version")),
                         str(a.get("checksum")), str(a.get("uri"))))
    elif key == "invalid_bay_table":
        rows = [("不连续格 ID", "说明")]
        for bid in (data or []):
            rows.append((str(bid), "不连续格：已排除出自然格级校核（ADR-0008），"
                                   "单元级校核照常"))
    else:
        return

    n_rows = max(1, len(rows))
    n_cols = max(1, len(rows[0])) if rows else 1
    table = doc.add_table(rows=n_rows, cols=n_cols)
    try:
        table.style = "Table Grid"
    except Exception:
        pass

    for r, row in enumerate(rows):
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(row[c]) if c < len(row) else "")
            try:
                run.font.size = Pt(9)
            except Exception:
                pass

    # 把表格元素移动到占位段落之后
    try:
        para._element.addnext(table._tbl)
    except Exception as exc:
        warnings.append("表格定位失败 %s: %s" % (key, exc))


def _scan_unresolved(path):
    """重新打开文档，统计仍残留的 {{...}} 占位符。"""
    try:
        doc = Document(path)
    except Exception:
        return []
    found = []
    for para in doc.paragraphs:
        for m in PLACEHOLDER_RE.finditer(para.text or ""):
            found.append(m.group(0))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for m in PLACEHOLDER_RE.finditer(para.text or ""):
                        found.append(m.group(0))
    return found
